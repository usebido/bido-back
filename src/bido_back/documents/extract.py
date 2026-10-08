"""Attachments -> model-ready text, every piece tagged with where it came from.

The quote-extraction agent reads this text and must cite, for every value it
extracts, the exact text and its locator. So every chunk carries a marker:

- PDF:        `[quote.pdf p.2]` before the page text, `[quote.pdf p.2 table 1]`
              before a table, which is rendered one row per line, cells joined by ` | `;
- Excel:      `[offer.xlsx Prices!C14] 0.45` for every non-empty cell, one sheet
              row per line;
- CSV:        `[offer.csv C14] 0.45`, the same without a sheet;
- plain text: `[notes.txt]` before the text;
- Word:       `[rfq.docx]` before the text, one paragraph per line, table
              rows rendered like PDF tables (cells joined by ` | `).

Images and scanned PDFs (no text layer) are not read: they come back as a
"needs vision" marker so the caller can list them (no OCR in the prototype).
Sizes are bounded (bytes, pages, cells, characters); whatever is cut is
replaced by a `truncated` marker, never dropped silently.

Supplier documents are untrusted. Nothing here interprets their content; it is
only turned into text for a model that treats it strictly as data.
"""

import csv
import email
import email.policy
import io
import logging
import re
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import PurePath
from typing import Literal, Protocol, cast
from xml.etree import ElementTree

logger = logging.getLogger(__name__)

DocumentKind = Literal["pdf", "sheet", "csv", "text", "docx", "image", "unsupported"]
SkipReason = Literal["too_large", "unreadable", "unsupported_format"]


@dataclass(frozen=True, slots=True)
class Limits:
    max_bytes: int = 25 * 1024 * 1024
    max_pdf_pages: int = 30
    # Non-empty cells per workbook or CSV file.
    max_cells: int = 5_000
    # Rows scanned per sheet, empty or not (guards against bogus sheet dimensions).
    max_rows: int = 10_000
    # Characters of model-ready text per document, markers included.
    max_chars: int = 120_000


DEFAULT_LIMITS = Limits()


@dataclass(frozen=True, slots=True)
class DocumentText:
    # The attachment's name as it appears in the markers (see `marker_name`).
    filename: str
    kind: DocumentKind
    # Model-ready text with locator markers. For documents that were not read
    # (images, scanned PDFs, unsupported or broken files) it is one marker line
    # saying so.
    text: str
    needs_vision: bool = False
    truncated: bool = False
    skipped: SkipReason | None = None

    @property
    def readable(self) -> bool:
        return not self.needs_vision and self.skipped is None


def marker_name(filename: str) -> str:
    """The attachment name as used inside `[...]` markers.

    Brackets and line breaks would break the markers, so they are replaced;
    the result is what a SourceRef's `attachment` must equal.
    """
    name = re.sub(r"[\[\]\r\n\t]+", "_", PurePath(filename.replace("\\", "/")).name).strip()
    return name[:120] or "attachment"


def extract_text(
    filename: str,
    content_type: str,
    data: bytes,
    *,
    limits: Limits = DEFAULT_LIMITS,
) -> DocumentText:
    """Turn one attachment into model-ready text. Never raises on bad input."""
    name = marker_name(filename)
    if len(data) > limits.max_bytes:
        return _skipped(name, "unsupported", "too_large", "file too large, not read")

    kind = _detect(name, content_type, data)
    try:
        match kind:
            case "pdf":
                return _pdf(name, data, limits)
            case "sheet":
                return _xlsx(name, data, limits)
            case "csv":
                return _csv(name, _decode(data), limits)
            case "text":
                return _text(name, content_type, data, limits)
            case "docx":
                return _docx(name, data, limits)
            case "image":
                return DocumentText(
                    filename=name,
                    kind="image",
                    text=f"[{name}] image attachment: needs vision, not read",
                    needs_vision=True,
                )
            case "unsupported":
                return _skipped(
                    name, "unsupported", "unsupported_format", "unsupported file type, not read"
                )
    except Exception as exc:  # noqa: BLE001 - a broken file must not fail the whole message
        # The type only: exception messages from parsers can quote file content.
        logger.warning("could not read attachment of kind %s: %s", kind, type(exc).__name__)
        return _skipped(name, kind, "unreadable", "file could not be read")


def render_documents(documents: Sequence[DocumentText]) -> str:
    """All documents as one text block, plus a list of what still needs vision."""
    parts = [doc.text for doc in documents]
    pending = [doc.filename for doc in documents if doc.needs_vision]
    if pending:
        parts.append("[not read: needs vision] " + ", ".join(pending))
    return "\n\n".join(parts)


# --- Detection ---------------------------------------------------------------

_PDF_TYPES = {"application/pdf", "application/x-pdf"}
_XLSX_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel.sheet.macroenabled.12",
}
_DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_CSV_TYPES = {"text/csv", "application/csv", "text/comma-separated-values"}
_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".tif", ".tiff", ".bmp", ".heic"}
_TEXT_SUFFIXES = {".txt", ".text", ".md", ".htm", ".html", ".eml"}
_TEXT_TYPES = {"message/rfc822"}


def _detect(name: str, content_type: str, data: bytes) -> DocumentKind:
    suffix = PurePath(name).suffix.lower()
    ctype = content_type.split(";", 1)[0].strip().lower()
    if data.startswith(b"%PDF-"):
        return "pdf"
    if suffix in (".xlsx", ".xlsm") or ctype in _XLSX_TYPES:
        return "sheet"
    if suffix == ".csv" or ctype in _CSV_TYPES:
        return "csv"
    if suffix == ".docx" or ctype == _DOCX_TYPE:
        return "docx"
    if suffix == ".pdf" or ctype in _PDF_TYPES:
        return "pdf"
    if suffix in _IMAGE_SUFFIXES or ctype.startswith("image/"):
        return "image"
    if suffix in _TEXT_SUFFIXES or ctype.startswith("text/") or ctype in _TEXT_TYPES:
        return "text"
    if data.startswith(b"PK\x03\x04") and b"xl/" in data[:4096]:
        return "sheet"
    if suffix == "" and _looks_like_text(data):
        return "text"
    return "unsupported"


def _looks_like_text(data: bytes) -> bool:
    head = data[:4096]
    if b"\x00" in head:
        return False
    try:
        head.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return True


def _skipped(name: str, kind: DocumentKind, reason: SkipReason, what: str) -> DocumentText:
    return DocumentText(filename=name, kind=kind, text=f"[{name}] {what}", skipped=reason)


class _Budget:
    """Caps the characters of one document; remembers whether it cut anything."""

    def __init__(self, max_chars: int) -> None:
        self._left = max_chars
        self.truncated = False

    def take(self, line: str) -> bool:
        if self._left - len(line) - 1 < 0:
            self.truncated = True
            return False
        self._left -= len(line) + 1
        return True


# --- PDF -----------------------------------------------------------------------


class _PdfTable(Protocol):
    @property
    def bbox(self) -> tuple[float, float, float, float]: ...

    def extract(self) -> list[list[str | None]]: ...


def _pdf(name: str, data: bytes, limits: Limits) -> DocumentText:
    import pdfplumber

    lines: list[str] = []
    budget = _Budget(limits.max_chars)
    pages_with_text = 0
    unread_visual_pages = False
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        total = len(pdf.pages)
        for number, page in enumerate(pdf.pages[: limits.max_pdf_pages], start=1):
            tables = cast(list[_PdfTable], page.find_tables())
            rest = page
            for table in tables:
                rest = rest.outside_bbox(table.bbox)
            text = (rest.extract_text() or "").strip()
            rendered = [_table_rows(table.extract()) for table in tables]
            if not text and not any(rendered):
                if page.chars or page.images or page.rects or page.curves or page.lines:
                    unread_visual_pages = True
                    lines.append(f"[{name} p.{number}] no text layer: needs vision, not read")
                else:
                    lines.append(f"[{name} p.{number}] empty page")
                continue
            pages_with_text += 1
            page_lines = [f"[{name} p.{number}]", *text.splitlines()] if text else []
            for index, rows in enumerate(rendered, start=1):
                if rows:
                    page_lines.append(f"[{name} p.{number} table {index}]")
                    page_lines.extend(rows)
            if not _take_all(lines, page_lines, budget):
                break

    if pages_with_text == 0:
        return DocumentText(
            filename=name,
            kind="pdf",
            text=f"[{name}] scanned PDF without a text layer: needs vision, not read",
            needs_vision=True,
        )
    # Mixed PDFs retain their readable pages, but cannot claim complete coverage.
    truncated = budget.truncated or total > limits.max_pdf_pages or unread_visual_pages
    if total > limits.max_pdf_pages:
        lines.append(f"[{name}] truncated: pages {limits.max_pdf_pages + 1}-{total} not read")
    elif budget.truncated:
        lines.append(f"[{name}] truncated: the rest of the document was not read")
    return DocumentText(filename=name, kind="pdf", text="\n".join(lines), truncated=truncated)


def _table_rows(table: list[list[str | None]]) -> list[str]:
    rows: list[str] = []
    for row in table:
        cells = [" ".join((cell or "").split()) for cell in row]
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


def _take_all(lines: list[str], new: Iterable[str], budget: _Budget) -> bool:
    for line in new:
        if not budget.take(line):
            return False
        lines.append(line)
    return True


# --- Excel ---------------------------------------------------------------------


class _Cell(Protocol):
    @property
    def value(self) -> object: ...


class _Sheet(Protocol):
    @property
    def title(self) -> str: ...

    def iter_rows(self, *, min_row: int, min_col: int) -> Iterable[tuple[_Cell, ...]]: ...


class _Workbook(Protocol):
    @property
    def worksheets(self) -> list[_Sheet]: ...

    def close(self) -> None: ...


def _load_workbook(data: bytes) -> _Workbook:
    # openpyxl ships no type information; contain it here.
    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # pyright: ignore[reportUnknownVariableType]
    return cast(_Workbook, workbook)


def _xlsx(name: str, data: bytes, limits: Limits) -> DocumentText:
    workbook = _load_workbook(data)
    lines: list[str] = []
    budget = _Budget(limits.max_chars)
    cells_left = limits.max_cells
    truncated = False
    try:
        for sheet in workbook.worksheets:
            prefix = f"{name} {_sheet_label(sheet.title)}!"
            rows = sheet.iter_rows(min_row=1, min_col=1)
            for row_number, row in enumerate(rows, start=1):
                if row_number > limits.max_rows:
                    lines.append(f"[{prefix[:-1]}] truncated: rows after {limits.max_rows}")
                    truncated = True
                    break
                cells = _row_cells(prefix, row_number, (c.value for c in row), _cell_hint(row))
                if not cells:
                    continue
                if len(cells) > cells_left or not budget.take(" | ".join(cells)):
                    truncated = True
                    break
                cells_left -= len(cells)
                lines.append(" | ".join(cells))
            if truncated:
                break
    finally:
        workbook.close()
    if truncated:
        lines.append(f"[{name}] truncated: the rest of the workbook was not read")
    if not lines:
        lines.append(f"[{name}] empty workbook")
    return DocumentText(filename=name, kind="sheet", text="\n".join(lines), truncated=truncated)


def _sheet_label(title: str) -> str:
    return re.sub(r"[\[\]\r\n\t]+", "_", title)


def _cell_hint(row: tuple[_Cell, ...]) -> list[str | None]:
    """Per cell: its number format when it names a currency (the value alone does not)."""
    hints: list[str | None] = []
    for cell in row:
        fmt = getattr(cell, "number_format", None)
        if isinstance(fmt, str) and any(mark in fmt for mark in ("€", "$", "£", "¥", "[$")):
            hints.append(fmt)
        else:
            hints.append(None)
    return hints


def _row_cells(
    prefix: str,
    row_number: int,
    values: Iterable[object],
    hints: Sequence[str | None] | None = None,
) -> list[str]:
    cells: list[str] = []
    for index, value in enumerate(values):
        text = format_cell(value)
        if not text:
            continue
        hint = hints[index] if hints is not None and index < len(hints) else None
        suffix = f" (format: {hint})" if hint else ""
        cells.append(f"[{prefix}{column_letter(index + 1)}{row_number}] {text}{suffix}")
    return cells


def column_letter(index: int) -> str:
    """1 -> A, 26 -> Z, 27 -> AA."""
    letters = ""
    while index > 0:
        index, rest = divmod(index - 1, 26)
        letters = chr(ord("A") + rest) + letters
    return letters


def format_cell(value: object) -> str:
    """A cell value as plain text: numbers without exponent or float noise."""
    match value:
        case None:
            return ""
        case bool():
            return "TRUE" if value else "FALSE"
        case int():
            return str(value)
        case float():
            if value.is_integer() and abs(value) < 1e15:
                return str(int(value))
            try:
                return format(Decimal(repr(value)), "f")
            except InvalidOperation:
                return repr(value)
        case Decimal():
            return format(value, "f")
        case datetime():
            if value.time() == time(0, 0):
                return value.date().isoformat()
            return value.isoformat(sep=" ")
        case date() | time():
            return value.isoformat()
        case str():
            return " ".join(value.split())
        case _:
            return " ".join(str(value).split())


# --- CSV -----------------------------------------------------------------------


def _csv(name: str, text: str, limits: Limits) -> DocumentText:
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
            text[:4096], delimiters=",;\t|"
        )
    except csv.Error:
        dialect = csv.excel
    lines: list[str] = []
    budget = _Budget(limits.max_chars)
    cells_left = limits.max_cells
    truncated = False
    for row_number, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        if row_number > limits.max_rows:
            truncated = True
            break
        cells = _row_cells(f"{name} ", row_number, row)
        if not cells:
            continue
        if len(cells) > cells_left or not budget.take(" | ".join(cells)):
            truncated = True
            break
        cells_left -= len(cells)
        lines.append(" | ".join(cells))
    if truncated:
        lines.append(f"[{name}] truncated: the rest of the file was not read")
    if not lines:
        lines.append(f"[{name}] empty file")
    return DocumentText(filename=name, kind="csv", text="\n".join(lines), truncated=truncated)


# --- Plain text ------------------------------------------------------------------


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


class _HtmlText(HTMLParser):
    _BLOCK = frozenset({"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "table"})
    _SKIP = frozenset({"script", "style", "head"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skipping += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")
        elif tag in ("td", "th"):
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skipping:
            self._skipping -= 1

    def handle_data(self, data: str) -> None:
        if not self._skipping:
            self.parts.append(data)


def html_to_text(html: str) -> str:
    parser = _HtmlText()
    parser.feed(html)
    parser.close()
    lines = (" ".join(line.split()) for line in "".join(parser.parts).splitlines())
    return "\n".join(line for line in lines if line)


def _text(name: str, content_type: str, data: bytes, limits: Limits) -> DocumentText:
    suffix = PurePath(name).suffix.lower()
    ctype = content_type.split(";", 1)[0].strip().lower()
    if suffix == ".eml" or ctype == "message/rfc822":
        body = _email_body(data)
    elif suffix in (".htm", ".html") or ctype == "text/html":
        body = html_to_text(_decode(data))
    else:
        body = _decode(data)
    body = body.replace("\r\n", "\n").strip()
    budget = _Budget(limits.max_chars)
    lines: list[str] = [f"[{name}]"]
    truncated = not _take_all(lines, body.splitlines(), budget)
    if truncated:
        lines.append(f"[{name}] truncated: the rest of the file was not read")
    return DocumentText(filename=name, kind="text", text="\n".join(lines), truncated=truncated)


def _email_body(data: bytes) -> str:
    """The text body of an attached email (a forwarded quote). Its own attachments are not read."""
    message = email.message_from_bytes(data, policy=email.policy.default)
    subject = str(message.get("subject", "") or "")
    part = message.get_body(preferencelist=("plain", "html"))
    content = ""
    if part is not None:
        payload: object = part.get_content()
        content = payload if isinstance(payload, str) else ""
        if part.get_content_subtype() == "html":
            content = html_to_text(content)
    return f"Subject: {subject}\n{content}" if subject else content


# --- Word (.docx) -----------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# The document part of a .docx is small XML; refuse a bomb before parsing it.
_MAX_DOCX_XML_BYTES = 20 * 1024 * 1024


def _docx(name: str, data: bytes, limits: Limits) -> DocumentText:
    """Paragraph and table text of `word/document.xml` (no headers, footers, images)."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        info = archive.getinfo("word/document.xml")
        if info.file_size > _MAX_DOCX_XML_BYTES:
            return _skipped(name, "docx", "too_large", "document too large, not read")
        # Expat (>= 2.4.1, bundled with Python 3.12) refuses entity-expansion
        # bombs and ElementTree never fetches external entities or DTDs; the
        # part's size is capped above.
        root = ElementTree.fromstring(archive.read(info))  # noqa: S314
    body = root.find(f"{_W}body")
    lines: list[str] = [f"[{name}]"]
    budget = _Budget(limits.max_chars)
    truncated = False
    for block in list(body) if body is not None else []:
        if block.tag == f"{_W}p":
            rendered = [_docx_text(block)]
        elif block.tag == f"{_W}tbl":
            rendered = _docx_table(block)
        else:
            continue
        if not _take_all(lines, (line for line in rendered if line), budget):
            truncated = True
            break
    if truncated:
        lines.append(f"[{name}] truncated: the rest of the document was not read")
    return DocumentText(filename=name, kind="docx", text="\n".join(lines), truncated=truncated)


def _docx_text(element: ElementTree.Element) -> str:
    return " ".join("".join(t.text or "" for t in element.iter(f"{_W}t")).split())


def _docx_table(table: ElementTree.Element) -> list[str]:
    rows: list[str] = []
    for row in table.iter(f"{_W}tr"):
        cells = [_docx_text(cell) for cell in row.iter(f"{_W}tc")]
        if any(cells):
            rows.append(" | ".join(cells))
    return rows
