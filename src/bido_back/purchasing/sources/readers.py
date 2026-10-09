"""File adapters into the same located evidence used by API/connector table adapters.

Workbooks (.xlsx, .xlsm, .xltx, .xltm) are read cell by cell without running macros or
formulas: a formula cell holds the value saved with the file (disclosed), or is empty
when no value was saved. Delimited exports (.csv, .tsv, and .txt files whose lines share
a tab, semicolon or pipe structure) become tables; their encoding falls back from UTF-8
to Windows-1252, disclosed. Other files go through the text extraction adapter.
"""

import codecs
import csv
import io
import re
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from typing import Protocol, cast

from bido_back.documents.extract import Limits, extract_text
from bido_back.purchasing.sources.contracts import EvidenceTable, Issue

# An assessment holds its sources and records in memory, about 2 KB per mapped cell: at
# these limits one call peaks near 500 MB (measured 2026-10-09 on a 218,000-cell export).
# Larger sources need a columnar store (docs/large-sources.md, step 3), not higher limits.
MAX_SOURCE_CELLS = 200_000
MAX_SOURCE_CHARS = 2_000_000
MAX_SOURCE_ROWS = 50_000
MAX_SOURCE_BYTES = 25 * 1024 * 1024

_WORKBOOKS = frozenset({"xlsx", "xlsm", "xltx", "xltm"})
_DELIMITED = frozenset({"csv", "tsv", "tab"})
_DELIMITED_TYPES = frozenset(
    {"text/csv", "application/csv", "text/comma-separated-values", "text/tab-separated-values"}
)
# Tie order: a decimal-comma export splits as consistently by "," as by ";".
_DELIMITERS = ("\t", ";", "|", ",")
# Plain text is a table only with a structure prose does not have.
_TEXT_DELIMITERS = ("\t", ";", "|")
_RULE = re.compile(r"[\s|+=_\-]*")
_SAMPLE_LINES = 200


@dataclass(frozen=True)
class SourceDocument:
    source_id: str
    filename: str
    content_type: str
    data: bytes


class _Cell(Protocol):
    @property
    def value(self) -> object: ...

    @property
    def number_format(self) -> str: ...

    @property
    def data_type(self) -> str: ...


class _Sheet(Protocol):
    @property
    def title(self) -> str: ...

    def iter_rows(self) -> Iterable[tuple[_Cell, ...]]: ...


class _Workbook(Protocol):
    @property
    def worksheets(self) -> list[_Sheet]: ...

    def close(self) -> None: ...


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def _xlsx(source: SourceDocument) -> tuple[list[EvidenceTable], list[Issue]]:
    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    # Macros are never loaded (keep_vba=False) and formulas are never evaluated.
    values = cast(_Workbook, load_workbook(io.BytesIO(source.data), read_only=True, data_only=True))
    formulas = cast(
        _Workbook, load_workbook(io.BytesIO(source.data), read_only=True, data_only=False)
    )
    tables: list[EvidenceTable] = []
    issues: list[Issue] = []
    cells, chars = 0, 0
    try:
        for sheet, raw_sheet in zip(values.worksheets, formulas.worksheets, strict=True):
            rows: list[list[str | None]] = []
            numbers: set[tuple[int, int]] = set()
            cached = missing = 0
            for number, (row, raw_row) in enumerate(
                zip(sheet.iter_rows(), raw_sheet.iter_rows(), strict=True), start=1
            ):
                cells += len(row)
                if number > MAX_SOURCE_ROWS or cells > MAX_SOURCE_CELLS:
                    raise ValueError("source_limit")
                rendered: list[str | None] = []
                for column, (cell, raw) in enumerate(zip(row, raw_row, strict=True), start=1):
                    value = cell.value
                    if raw.data_type == "f":
                        # The value saved with the file; an error or no saved value is empty.
                        if value is None or cell.data_type == "e":
                            missing += 1
                            value = None
                        else:
                            cached += 1
                    elif cell.data_type == "e":
                        value = None
                    text = _text(value)
                    if isinstance(value, int | float) and not isinstance(value, bool):
                        numbers.add((number, column))
                        if isinstance(value, int) and re.fullmatch(r"0+", cell.number_format):
                            text = str(value).zfill(len(cell.number_format))
                    rendered.append(text)
                    chars += len(text or "")
                    if chars > MAX_SOURCE_CHARS:
                        raise ValueError("source_limit")
                rows.append(rendered)
            tables.append(
                EvidenceTable(
                    source_id=source.source_id, name=sheet.title, rows=rows, numbers=numbers
                )
            )
            if cached:
                issues.append(
                    Issue(
                        code="formula_values",
                        severity="warning",
                        message=f"{sheet.title}: {cached} formula cell(s) were read as the "
                        "values saved with the file; formulas were not executed.",
                    )
                )
            if missing:
                issues.append(
                    Issue(
                        code="formula_without_value",
                        severity="warning",
                        message=f"{sheet.title}: {missing} formula cell(s) have no saved value "
                        "or hold an error and were read as empty; nothing was computed.",
                    )
                )
    finally:
        values.close()
        formulas.close()
    return tables, issues


def _decode(data: bytes) -> tuple[str, str | None]:
    """The text and, when it is not UTF-8/UTF-16, the fallback encoding that was used."""
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16"), None
    try:
        return data.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1252"), "Windows-1252"
    except UnicodeDecodeError:
        return data.decode("latin-1"), "Latin-1"


def _split(text: str, delimiter: str) -> list[list[str]] | None:
    try:
        return list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except csv.Error:
        return None


def _consistency(rows: list[list[str]]) -> tuple[float, int, int]:
    """Share of non-empty rows with the most common width, that width, and its row count."""
    widths = [len(row) for row in rows if any(cell.strip() for cell in row)]
    if not widths:
        return 0.0, 0, 0
    width, count = Counter(widths).most_common(1)[0]
    return count / len(widths), width, count


def _delimiter(text: str, candidates: tuple[str, ...]) -> str | None:
    sample = "\n".join(text.splitlines()[:_SAMPLE_LINES])
    best: tuple[float, int] | None = None
    chosen: str | None = None
    for delimiter in candidates:
        rows = _split(sample, delimiter)
        if rows is None:
            continue
        share, width, count = _consistency(rows)
        if width < 2 or count < 2:
            continue
        if best is None or (share, width) > best:
            best, chosen = (share, width), delimiter
    if best is None or best[0] < 0.8:
        return None
    return chosen


def _framed(row: list[str]) -> bool:
    return not row[0].strip() and not row[-1].strip()


def _cells(rows: list[list[str]], delimiter: str) -> list[list[str | None]]:
    if delimiter == "|":
        # Fixed-width list exports frame cells with pipes and draw rules with dashes.
        rows = [[] if _RULE.fullmatch("|".join(row)) else row for row in rows]
        wide = [row for row in rows if len(row) >= 3]
        if wide and sum(1 for row in wide if _framed(row)) >= 0.8 * len(wide):
            rows = [row[1:-1] if len(row) >= 3 and _framed(row) else row for row in rows]
        rows = [[cell.strip() for cell in row] for row in rows]
    return [[cell if cell.strip() else None for cell in row] for row in rows]


def _delimited(
    source: SourceDocument, *, text_only: bool
) -> tuple[list[EvidenceTable], list[Issue]] | None:
    """A delimited export as one table; None when a text file has no table structure."""
    text, fallback = _decode(source.data)
    delimiter = _delimiter(text, _TEXT_DELIMITERS if text_only else _DELIMITERS)
    if delimiter is None and text_only:
        return None
    rows = _split(text, delimiter or ",")
    if rows is None:
        if text_only:
            return None
        raise ValueError("unreadable")
    table = _cells(rows, delimiter or ",")
    if (
        len(table) > MAX_SOURCE_ROWS
        or sum(len(row) for row in table) > MAX_SOURCE_CELLS
        or sum(len(value or "") for row in table for value in row) > MAX_SOURCE_CHARS
    ):
        return [], [_unread(source, "source_limit")]
    issues: list[Issue] = []
    if fallback is not None:
        issues.append(
            Issue(
                code="source_encoding",
                severity="warning",
                message=f"Source {source.source_id} is not UTF-8 and was decoded as {fallback}; "
                "accented characters depend on that encoding.",
            )
        )
    return [EvidenceTable(source_id=source.source_id, name=source.filename, rows=table)], issues


def read_source(source: SourceDocument) -> tuple[list[EvidenceTable], list[Issue]]:
    """Read every sheet or the text adapter's disclosed coverage; never run source code."""
    if len(source.data) > MAX_SOURCE_BYTES:
        return [], [_unread(source, "source_limit")]
    try:
        suffix = source.filename.lower().rsplit(".", 1)[-1] if "." in source.filename else ""
        content_type = source.content_type.split(";", 1)[0].strip().lower()
        if suffix in _WORKBOOKS:
            return _xlsx(source)
        if suffix in _DELIMITED or (suffix == "" and content_type in _DELIMITED_TYPES):
            found = _delimited(source, text_only=False)
            if found is not None:
                return found
        if suffix == "txt":
            found = _delimited(source, text_only=True)
            if found is not None:
                return found
        document = extract_text(
            source.filename,
            source.content_type,
            source.data,
            limits=Limits(
                max_chars=MAX_SOURCE_CHARS,
                max_cells=MAX_SOURCE_CELLS,
                max_rows=MAX_SOURCE_ROWS,
            ),
        )
        if not document.readable:
            return [], [_unread(source, "needs_vision" if document.needs_vision else "unreadable")]
        text_rows: list[list[str | None]] = [[line] for line in document.text.splitlines()]
        if len(text_rows) > MAX_SOURCE_ROWS:
            return [], [_unread(source, "source_limit")]
        issues = [_unread(source, "partial_coverage")] if document.truncated else []
        return [
            EvidenceTable(
                source_id=source.source_id,
                name=source.filename,
                layout="text",
                rows=text_rows,
            )
        ], issues
    except ValueError:
        return [], [_unread(source, "unreadable_or_limit")]
    except Exception:
        # Parser exceptions may contain customer content; expose a stable code only.
        return [], [_unread(source, "unreadable")]


def _unread(source: SourceDocument, code: str) -> Issue:
    return Issue(code=code, severity="blocker", message=f"Source {source.source_id}: {code}.")
