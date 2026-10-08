"""Bounded, lossless cell windows for spreadsheet understanding before import."""

import io
from collections.abc import Iterable
from datetime import date, datetime, time
from typing import Literal, Protocol, cast

from pydantic import BaseModel, Field

from bido_back.agents.catalog.contracts import ImportFieldCatalog
from bido_back.purchasing.importers import SPECS, ImportFileError, read_table
from bido_back.purchasing.workbook import WorkbookFileError


class SpreadsheetWindow(BaseModel):
    sheet: str | None = Field(
        default=None, description="Null lists all sheets without reading cells"
    )
    first_row: int = Field(default=1, ge=1, le=1_048_576)
    first_column: int = Field(default=1, ge=1, le=16_384)
    row_count: int = Field(default=50, ge=1, le=100)
    column_count: int = Field(default=20, ge=1, le=50)


class SheetInfo(BaseModel):
    name: str
    rows: int
    columns: int
    state: str


class SpreadsheetCell(BaseModel):
    row: int
    column: int
    value: str | None
    kind: Literal["text", "number", "boolean", "date", "error", "empty"]
    formula: str | None = None
    number_format: str | None = None


class SpreadsheetInspection(BaseModel):
    filename: str
    sheets: list[SheetInfo]
    import_catalog: list[ImportFieldCatalog] = Field(default_factory=list[ImportFieldCatalog])
    sheet: str | None = None
    first_row: int | None = None
    last_row: int | None = None
    first_column: int | None = None
    last_column: int | None = None
    next_row: int | None = None
    next_column: int | None = None
    cells: list[SpreadsheetCell] = Field(default_factory=list[SpreadsheetCell])


class _Cell(Protocol):
    @property
    def value(self) -> object: ...

    @property
    def data_type(self) -> str: ...

    @property
    def number_format(self) -> str: ...


class _Sheet(Protocol):
    title: str
    max_row: int | None
    max_column: int | None
    sheet_state: str

    def iter_rows(
        self, *, min_row: int, max_row: int, min_col: int, max_col: int
    ) -> Iterable[tuple[_Cell, ...]]: ...


class _Book(Protocol):
    worksheets: list[_Sheet]

    def close(self) -> None: ...


def _cell(row: int, column: int, value: object, *, error: bool = False) -> SpreadsheetCell:
    kind: Literal["text", "number", "boolean", "date", "error", "empty"] = "text"
    text = str(value) if value is not None else None
    if error:
        kind = "error"
    elif value is None:
        kind = "empty"
    elif isinstance(value, bool):
        kind, text = "boolean", str(value).lower()
    elif isinstance(value, datetime | date | time):
        kind, text = "date", value.isoformat()
    elif isinstance(value, int | float):
        kind = "number"
    return SpreadsheetCell(row=row, column=column, value=text, kind=kind)


def _window(
    filename: str, sheets: list[SheetInfo], request: SpreadsheetWindow
) -> SpreadsheetInspection:
    result = SpreadsheetInspection(
        filename=filename,
        sheets=sheets,
        sheet=request.sheet,
        import_catalog=[
            ImportFieldCatalog(
                import_type=kind,
                fields=[s.name for s in specs],
                required=[s.name for s in specs if s.required],
            )
            for kind, specs in SPECS.items()
        ],
    )
    if request.sheet is None:
        return result
    sheet = next((s for s in sheets if s.name == request.sheet), None)
    if sheet is None:
        raise WorkbookFileError("sheet not found; list sheets first")
    if request.first_row > sheet.rows or request.first_column > sheet.columns:
        raise WorkbookFileError("requested window is outside the sheet")
    result.first_row = request.first_row
    result.last_row = min(sheet.rows, request.first_row + request.row_count - 1)
    result.first_column = request.first_column
    result.last_column = min(sheet.columns, request.first_column + request.column_count - 1)
    result.next_row = result.last_row + 1 if result.last_row < sheet.rows else None
    result.next_column = result.last_column + 1 if result.last_column < sheet.columns else None
    return result


def inspect_spreadsheet(
    data: bytes, filename: str, request: SpreadsheetWindow
) -> SpreadsheetInspection:
    """List every sheet or read a rectangular window, never truncating cell contents.

    Formulas are returned as data alongside cached values, never evaluated.
    Empty cells are omitted; their coordinates remain implicit in the window.
    """
    if filename.lower().endswith(".csv"):
        try:
            rows = read_table(data, filename)
        except ImportFileError as exc:
            raise WorkbookFileError(str(exc)) from exc
        result = _window(
            filename,
            [
                SheetInfo(
                    name=filename,
                    rows=len(rows),
                    columns=max(map(len, rows), default=0),
                    state="visible",
                )
            ],
            request,
        )
        if request.sheet is not None:
            result.cells = [
                _cell(r, c, value)
                for r, row in enumerate(rows, start=1)
                if request.first_row <= r < request.first_row + request.row_count
                for c, value in enumerate(row, start=1)
                if request.first_column <= c < request.first_column + request.column_count
                and value is not None
                and value != ""
            ]
    elif filename.lower().endswith(".xlsx"):
        result = _inspect_xlsx(data, filename, request)
    else:
        raise WorkbookFileError("spreadsheets must be .xlsx or .csv")
    if sum(len(c.value or "") + len(c.formula or "") for c in result.cells) > 100_000:
        raise WorkbookFileError("window is too large; request fewer rows or columns")
    return result


def _inspect_xlsx(data: bytes, filename: str, request: SpreadsheetWindow) -> SpreadsheetInspection:
    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    books: list[_Book] = []
    try:
        book = cast(_Book, load_workbook(io.BytesIO(data), read_only=True, data_only=False))
        books.append(book)
        result = _window(
            filename,
            [
                SheetInfo(
                    name=s.title,
                    rows=s.max_row or 0,
                    columns=s.max_column or 0,
                    state=s.sheet_state,
                )
                for s in book.worksheets
            ],
            request,
        )
        if request.sheet is None:
            return result
        cached = cast(_Book, load_workbook(io.BytesIO(data), read_only=True, data_only=True))
        books.append(cached)
        sheet = next(s for s in book.worksheets if s.title == request.sheet)
        values = next(s for s in cached.worksheets if s.title == request.sheet)
        bounds = dict(
            min_row=request.first_row,
            max_row=result.last_row or request.first_row,
            min_col=request.first_column,
            max_col=result.last_column or request.first_column,
        )
        for row, (raw, saved) in enumerate(
            zip(sheet.iter_rows(**bounds), values.iter_rows(**bounds), strict=True),
            start=request.first_row,
        ):
            for column, (cell, value) in enumerate(
                zip(raw, saved, strict=True), start=request.first_column
            ):
                if cell.value is None:
                    continue
                item = _cell(row, column, value.value, error=value.data_type == "e")
                if cell.data_type == "f":
                    item.formula = str(cell.value)
                item.number_format = cell.number_format
                result.cells.append(item)
        return result
    except WorkbookFileError:
        raise
    except Exception as exc:
        raise WorkbookFileError("the file is not a readable .xlsx workbook") from exc
    finally:
        for book in books:
            book.close()
