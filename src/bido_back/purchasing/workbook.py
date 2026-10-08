"""Importing whatever spreadsheet the buyer has: every sheet, any layout.

Buyers send what their ERP or their own Excel produces. There is no template:

1. Read every sheet (`read_workbook`).
2. Find the sections (`map_workbook`). Rules first: in the first rows of each
   sheet, a row whose headers name every required field of an import type is
   a header row (Portuguese, English and Italian synonyms). A sheet the rules
   cannot read goes to the sheet mapper (Luna), which only proposes WHERE the
   table is and WHAT each column means; one sheet may hold two sections (e.g.
   purchase history plus the supplier's email). Its proposal is checked:
   known sheet, known type, fields of that type, columns inside the sheet, no
   column used twice.
3. Validate every value with the importers (`parse_mapped`), as for any import.
4. Ask before writing when unsure: a section with low confidence, a missing
   required field, a question from the mapper, or more than a fifth of its rows
   failing makes the import `needs_confirmation`. Nothing is written until
   the buyer confirms; then the valid rows are imported, in dependency order.

The model never sees more than a preview (first rows, cells cut short) and
never produces a value that is written.

What ERP exports (SAP and the like) need, and how it is read:

* Numbers written as text. A number such as "888,000" reads both ways; it is
  resolved with the numbers of its own column, then of its section, then of the
  whole workbook ("1.850,000" or "14,73" elsewhere settle the comma), and as a
  last hint with day-first dotted dates ("08.10.2026"). With no evidence the
  row is reported, never guessed (`importers.decimal_evidence`).
* Totals and subtotals under a table ("Total", "*") are skipped.
* A multi-level BOM (a level column: "0", ".1", "..2") is exploded to the parts
  bought: each leaf's quantity is multiplied through its assemblies and divided
  by the base quantity of the level-0 row; assemblies are not imported (they are
  built, not bought, and their stock is not netted). The product is the level-0
  row, else the product named in the title rows ("Material: MT-120-ASM").
* Supplier codes: a sheet listing suppliers by code only ("Fornecedor 100231")
  takes their names from a sheet that has both (purchase orders).
* Stock: only unrestricted stock counts as on hand; quality-inspection and
  blocked stock cannot go into a build and are left out (the report says so).
* Currencies stay as written per purchase line; the analysis handles them.
"""

import io
import logging
import re
from collections.abc import Awaitable, Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Final, Literal, Protocol, cast
from uuid import UUID, uuid4

from pydantic import BaseModel, Field
from sqlalchemy import exists, select
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession, async_sessionmaker

from bido_back.agents.catalog.contracts import (
    ImportFieldCatalog,
    SheetMappingInput,
    SheetMappingOutput,
    SheetPreview,
)
from bido_back.agents.config import RunnerConfig
from bido_back.agents.context import ToolContext
from bido_back.auth import Principal
from bido_back.db import tenant_session
from bido_back.db.models import Contact, Supplier, SupplierContact, TemplateUpload
from bido_back.logs import safe_reason
from bido_back.purchasing.importers import (
    MAX_ROWS,
    SPECS,
    ImportFileError,
    ImportType,
    ParsedTable,
    RowError,
    decimal_evidence,
    map_columns,
    norm,
    parse_mapped,
    read_table,
)
from bido_back.quotes.parsing import DecimalSeparator
from bido_back.storage import BlobStorage

logger = logging.getLogger(__name__)

Db = AsyncEngine | async_sessionmaker[AsyncSession]

SheetMapper = Callable[[SheetMappingInput, ToolContext], Awaitable[SheetMappingOutput]]

HEADER_SCAN_ROWS: Final = 30
PREVIEW_ROWS: Final = 25
PREVIEW_COLUMNS: Final = 30
PREVIEW_CELL_CHARS: Final = 40
MAX_SHEETS: Final = 20
CONFIDENT: Final = 0.7
MAX_FAILED_SHARE: Final = 0.2
# Inventory can supply item descriptions and units missing from order history.
_ORDER: Final[tuple[ImportType, ...]] = (
    "items",
    "bom",
    "supplier_contacts",
    "inventory",
    "purchase_history",
    "production_plan",
)
_DETAIL_TYPES: Final[tuple[ImportType, ...]] = (
    "bom",
    "purchase_history",
    "inventory",
    "production_plan",
)


@dataclass(frozen=True, slots=True)
class SheetGrid:
    name: str
    rows: list[list[object]]

    @property
    def width(self) -> int:
        return max((len(r) for r in self.rows), default=0)


@dataclass(slots=True)
class Section:
    sheet: str
    import_type: ImportType
    header_row: int  # 0-based index into the sheet's rows
    first_row: int  # 0-based, inclusive
    last_row: int | None  # 0-based, inclusive; None = to the end
    index: dict[str, int]
    confidence: float
    source: Literal["rules", "model"]
    question: str | None = None


class WorkbookFileError(Exception):
    pass


# --- Reading -------------------------------------------------------------------------------


def read_workbook(data: bytes, filename: str) -> list[SheetGrid]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        return [SheetGrid(name=filename, rows=read_table(data, filename))]
    if lower.endswith(".xls"):
        raise WorkbookFileError("old .xls files are not supported; save it as .xlsx or .csv")
    if not lower.endswith(".xlsx"):
        raise WorkbookFileError("spreadsheets must be .xlsx or .csv")
    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    try:
        workbook = cast(_Workbook, load_workbook(io.BytesIO(data), read_only=True, data_only=True))
    except Exception as exc:  # openpyxl raises several unrelated types on bad input
        raise WorkbookFileError("the file is not a readable .xlsx workbook") from exc
    sheets: list[SheetGrid] = []
    try:
        if len(workbook.worksheets) > MAX_SHEETS:
            raise WorkbookFileError(
                f"imports are limited to {MAX_SHEETS} sheets; split the workbook"
            )
        for sheet in workbook.worksheets:
            rows: list[list[object]] = []
            for raw in sheet.iter_rows(values_only=True):
                rows.append(list(raw))
                if len(rows) > MAX_ROWS + HEADER_SCAN_ROWS:
                    raise WorkbookFileError(f"sheets are limited to {MAX_ROWS} rows")
            sheets.append(SheetGrid(name=sheet.title, rows=rows))
    finally:
        workbook.close()
    return sheets


class _Sheet(Protocol):
    @property
    def title(self) -> str: ...

    def iter_rows(self, *, values_only: bool) -> Iterable[tuple[object, ...]]: ...


class _Workbook(Protocol):
    @property
    def worksheets(self) -> list[_Sheet]: ...

    def close(self) -> None: ...


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _cell(value: object) -> str:
    text = " ".join(str(value).split())
    return text[:PREVIEW_CELL_CHARS] + ("…" if len(text) > PREVIEW_CELL_CHARS else "")


def preview(grid: SheetGrid) -> str:
    lines: list[str] = []
    for number, row in enumerate(grid.rows[:PREVIEW_ROWS], start=1):
        cells = [f"[{i}] {_cell(v)}" for i, v in enumerate(row[:PREVIEW_COLUMNS]) if not _blank(v)]
        if cells:
            lines.append(f"row {number}: " + " | ".join(cells))
    return "\n".join(lines)


# --- Finding the sections ----------------------------------------------------------------


def _rule_candidates(grid: SheetGrid) -> list[Section]:
    best: dict[ImportType, tuple[int, int, dict[str, int]]] = {}
    for at, row in enumerate(grid.rows[:HEADER_SCAN_ROWS]):
        if sum(1 for c in row if not _blank(c)) < 2:
            continue
        for import_type in SPECS:
            index, _, missing, _ = map_columns(import_type, row)
            if missing == ["product"] and import_type == "bom":
                # A BOM explosion names its product above the table or in its level-0 row.
                if "level" in index or _title_product(grid.rows[:at]) is not None:
                    missing = []
            if missing:
                continue
            score = len(index)
            if import_type not in best or score > best[import_type][1]:
                best[import_type] = (at, score, index)
    sections = [
        Section(
            sheet=grid.name,
            import_type=t,
            header_row=at,
            first_row=at + 1,
            last_row=None,
            index=index,
            confidence=0.9,
            source="rules",
        )
        for t, (at, _, index) in best.items()
    ]
    if any(s.import_type in _DETAIL_TYPES for s in sections):
        # Purchase lines and BOMs carry descriptions too: not a separate catalog.
        sections = [s for s in sections if s.import_type != "items"]
    if any(s.import_type in ("bom", "purchase_history") for s in sections):
        # A BOM's or a purchase line's quantity and date also read as a production
        # plan; the BOM or the purchase history wins.
        sections = [s for s in sections if s.import_type != "production_plan"]
    return sections


# "Material: MT-120-ASM   Centro: 1100", "Product MT-120", "Articolo: X-1". The code
# must hold a digit, so "Produto acabado: ..." does not read "acabado" as a code.
_TITLE_PRODUCT: Final = re.compile(
    r"\b(?:material|produto|product|prodotto|articolo|assembly|montagem|modelo|model|"
    r"conjunto|item)\b\s*[:=]?\s*([A-Za-z0-9][A-Za-z0-9._/-]*)",
    re.IGNORECASE,
)


def _title_product(rows: Sequence[Sequence[object]]) -> str | None:
    for row in rows:
        for cell in row:
            if not isinstance(cell, str):
                continue
            for match in _TITLE_PRODUCT.finditer(cell):
                code = match.group(1).rstrip(".")
                if any(c.isdigit() for c in code):
                    return code
    return None


def _catalog() -> list[ImportFieldCatalog]:
    return [
        ImportFieldCatalog(
            import_type=t,
            fields=[s.name for s in specs],
            required=[s.name for s in specs if s.required],
        )
        for t, specs in SPECS.items()
    ]


def _from_model(
    output: SheetMappingOutput, grids: Mapping[str, SheetGrid]
) -> tuple[list[Section], list[str]]:
    """Checked sections from the mapper's proposal, and the problems found."""
    sections: list[Section] = []
    problems: list[str] = []
    for proposed in output.sections:
        grid = grids.get(proposed.sheet)
        if grid is None:
            problems.append("mapping references an unknown sheet")
            continue
        if proposed.import_type == "ignore":
            continue
        import_type = proposed.import_type
        names = {s.name for s in SPECS[import_type]}
        header = proposed.header_row - 1
        first = proposed.first_data_row - 1
        last = proposed.last_data_row - 1 if proposed.last_data_row is not None else None
        if (
            not 0 <= header < len(grid.rows)
            or first <= header
            or first >= len(grid.rows)
            or (last is not None and not first <= last < len(grid.rows))
        ):
            problems.append(f"{proposed.sheet}: rows out of range")
            continue
        index: dict[str, int] = {}
        used: set[int] = set()
        bad = False
        for assignment in proposed.columns:
            if (
                assignment.field not in names
                or assignment.field in index
                or not 0 <= assignment.column < grid.width
                or assignment.column in used
            ):
                bad = True
                break
            index[assignment.field] = assignment.column
            used.add(assignment.column)
        if bad:
            problems.append(f"{proposed.sheet}: invalid column mapping for {import_type}")
            continue
        if any(
            s.sheet == grid.name
            and s.import_type == import_type
            and not used.isdisjoint(s.index.values())
            and first <= (s.last_row if s.last_row is not None else len(grid.rows) - 1)
            and s.first_row <= (last if last is not None else len(grid.rows) - 1)
            for s in sections
        ):
            problems.append(f"{proposed.sheet}: overlapping sections for {import_type}")
            continue
        sections.append(
            Section(
                sheet=grid.name,
                import_type=import_type,
                header_row=header,
                first_row=first,
                last_row=last,
                index=index,
                confidence=proposed.confidence,
                source="model",
                question=proposed.question,
            )
        )
    return sections, problems


async def map_workbook(
    grids: Sequence[SheetGrid],
    filename: str,
    *,
    mapper: SheetMapper | None,
    context: ToolContext,
) -> tuple[list[Section], list[str]]:
    """The sections to import: rules where they suffice, the mapper for the rest.

    A sheet where the rules only see a parts list or supplier contacts, with
    several columns left unexplained, also goes to the mapper (a purchase table
    with headers the rules do not know reads as "part + description" alone); its
    rule sections stay only if the mapper finds nothing there.
    """
    sections: list[Section] = []
    unresolved: list[SheetGrid] = []
    partial: dict[str, list[Section]] = {}
    for grid in grids:
        found = _rule_candidates(grid)
        if found and _partly_read(found, grid):
            partial[grid.name] = found
            unresolved.append(grid)
        elif found:
            sections.extend(found)
        elif any(not _blank(c) for row in grid.rows[:PREVIEW_ROWS] for c in row):
            unresolved.append(grid)
    problems: list[str] = []
    if unresolved:
        if mapper is None:
            problems.extend(
                f"{g.name}: could not find the table" for g in unresolved if g.name not in partial
            )
            sections.extend(s for found in partial.values() for s in found)
        else:
            output = await mapper(
                SheetMappingInput(
                    filename=filename,
                    sheets=[
                        SheetPreview(
                            sheet=g.name, rows=len(g.rows), columns=g.width, preview=preview(g)
                        )
                        for g in unresolved
                    ],
                    catalog=_catalog(),
                ),
                context,
            )
            mapped, problems = _from_model(output, {g.name: g for g in unresolved})
            sections.extend(mapped)
            read = {s.sheet for s in mapped}
            sections.extend(s for name, found in partial.items() if name not in read for s in found)
    return sections, problems


_WEAK_TYPES: Final[frozenset[ImportType]] = frozenset({"items", "supplier_contacts"})
_UNEXPLAINED_COLUMNS: Final = 3


def _partly_read(found: Sequence[Section], grid: SheetGrid) -> bool:
    if any(s.import_type not in _WEAK_TYPES for s in found):
        return False
    used = {i for s in found for i in s.index.values()}
    rows = {s.header_row for s in found}
    unexplained = {
        i
        for at in rows
        for i, cell in enumerate(grid.rows[at])
        if i not in used and not _blank(cell)
    }
    return len(unexplained) >= _UNEXPLAINED_COLUMNS


# --- Validating and importing ---------------------------------------------------------------

MAX_SUPPLIERS_WITHOUT_EMAIL: Final = 50


class SectionReport(BaseModel):
    sheet: str
    type: ImportType
    source: Literal["rules", "model"]
    header_row: int  # 1-based, as in the spreadsheet
    columns: dict[str, str]
    unknown_columns: list[str] = Field(default_factory=list[str])
    missing_required: list[str]
    rows_total: int
    rows_ok: int
    rows_failed: int
    # Totals/subtotals, and the assemblies of an exploded multi-level BOM.
    rows_skipped: int = 0
    errors: list[str]
    confidence: float
    question: str | None
    created: dict[str, int] = {}
    updated: dict[str, int] = {}
    skipped: dict[str, int] = {}
    # How the section was read, for the buyer (BOM exploded, stock left out...).
    notes: list[str] = []


class WorkbookImportReport(BaseModel):
    upload_id: UUID
    filename: str
    status: Literal["imported", "needs_confirmation", "nothing_found"]
    sections: list[SectionReport]
    # What to ask the buyer before importing; empty when imported.
    questions: list[str]
    problems: list[str]
    unimported_sheets: list[str] = Field(
        default_factory=list[str],
        description="Nonempty sheets with no parsed section; their data was not imported",
    )
    suppliers_without_email: list[str] = Field(
        default_factory=list[str],
        description=(
            "After an import: the buyer's suppliers that still have no contact email, by "
            f"name (at most {MAX_SUPPLIERS_WITHOUT_EMAIL}). They cannot be invited to an RFQ "
            "until one is added: tell the buyer now, not when the RFQ is created"
        ),
    )


@dataclass(slots=True)
class _Parsed:
    section: Section
    table: ParsedTable
    notes: list[str] = field(default_factory=list[str])
    report: SectionReport = field(init=False)

    def __post_init__(self) -> None:
        self.refresh()

    def refresh(self) -> None:
        t = self.table
        failed = t.total_rows - len(t.rows)
        self.report = SectionReport(
            sheet=self.section.sheet,
            type=self.section.import_type,
            source=self.section.source,
            header_row=self.section.header_row + 1,
            columns=t.columns,
            unknown_columns=t.unknown_columns,
            missing_required=t.missing_required,
            rows_total=t.total_rows,
            rows_ok=len(t.rows),
            rows_failed=failed,
            rows_skipped=t.skipped_rows,
            errors=[
                f"row {e.row}: {e.field + ': ' if e.field else ''}{e.message}"
                for e in t.errors[:10]
            ],
            confidence=round(self.section.confidence, 2),
            question=self.section.question,
            notes=list(self.notes),
        )

    def doubts(self) -> list[str]:
        r = self.report
        found: list[str] = []
        if r.question:
            found.append(r.question)
        elif r.confidence < CONFIDENT:
            found.append(f"Sheet '{r.sheet}': is this {r.type.replace('_', ' ')}?")
        if r.missing_required:
            found.append(
                f"Sheet '{r.sheet}' ({r.type}): which columns hold {', '.join(r.missing_required)}?"
            )
        if r.rows_total and r.rows_failed / r.rows_total > MAX_FAILED_SHARE:
            found.append(
                f"Sheet '{r.sheet}': {r.rows_failed} of {r.rows_total} rows have invalid "
                "values; import only the valid rows?"
            )
        return found


def parse_section(
    section: Section, grid: SheetGrid, decimal: DecimalSeparator | None
) -> ParsedTable:
    return _parse_section(section, grid, decimal)[0]


def _parse_section(
    section: Section, grid: SheetGrid, decimal: DecimalSeparator | None
) -> tuple[ParsedTable, list[str]]:
    """The section's validated rows and notes on how it was read."""
    header: list[object] = list(grid.rows[section.header_row])
    end = len(grid.rows) if section.last_row is None else section.last_row + 1
    rows: list[tuple[int, Sequence[object]]] = [
        (number + 1, grid.rows[number])
        for number in range(section.first_row, min(end, len(grid.rows)))
    ]
    index = dict(section.index)
    notes: list[str] = []
    if section.import_type == "bom" and "product" not in index:
        product = _bom_product(index, rows) or _title_product(grid.rows[: section.header_row])
        if product is not None:
            # The product as one more column, so every row is validated as usual.
            column = max(grid.width, len(header))
            header = [*header, *([None] * (column - len(header))), f"(product {product})"]
            rows = [(n, [*r, *([None] * (column - len(r))), product]) for n, r in rows]
            index["product"] = column
            notes.append(f"Product {product} taken from the sheet's title / level-0 row.")
    table = parse_mapped(section.import_type, header, index, rows, decimal=decimal)
    if section.import_type == "bom" and "level" in index and not table.missing_required:
        notes.extend(_explode_levels(table, index["level"], rows))
    return table, notes


def _level(cell: object) -> int | None:
    """A BOM level: 0, ".1", "..2" (SAP), 1, or outline numbering ("1.2.1" is 3)."""
    if isinstance(cell, bool) or cell is None:
        return None
    if isinstance(cell, int):
        return cell if cell >= 0 else None
    if isinstance(cell, float):
        return int(cell) if cell.is_integer() and cell >= 0 else None
    text = str(cell).strip()
    if (m := re.fullmatch(r"\.*(\d{1,2})", text)) is not None:
        return int(m.group(1))
    if re.fullmatch(r"\d{1,2}(?:\.\d{1,3})+", text) is not None:
        return text.count(".") + 1
    return None


def _bom_product(
    index: Mapping[str, int], rows: Sequence[tuple[int, Sequence[object]]]
) -> str | None:
    """The part number of the first level-0 row, if the BOM has levels."""
    level_at, part_at = index.get("level"), index.get("part_number")
    if level_at is None or part_at is None:
        return None
    for _, raw in rows:
        if level_at < len(raw) and part_at < len(raw) and _level(raw[level_at]) == 0:
            part = raw[part_at]
            if not _blank(part):
                return str(part).strip()
    return None


def _explode_levels(
    table: ParsedTable, level_at: int, rows: Sequence[tuple[int, Sequence[object]]]
) -> list[str]:
    """Replace a multi-level BOM's rows with its leaves, quantities per one product.

    A row is an assembly when the next row is one level deeper. A leaf's
    quantity is its own times every assembly's above it, divided by the base
    quantity of its level-0 row (the product; 1 when there is none). A leaf
    under an invalid assembly is reported, not imported.
    """
    parsed = dict(table.rows)
    failed = {e.row for e in table.errors}
    entries: list[tuple[int, int, dict[str, object] | None]] = []
    for offset, raw in rows:
        if offset not in parsed and offset not in failed:
            continue  # blank, total
        level = _level(raw[level_at] if level_at < len(raw) else None)
        if level is None:
            if offset in parsed:
                del parsed[offset]
                table.errors.append(RowError(offset, "level", "not a BOM level"))
            continue
        entries.append((offset, level, parsed.get(offset)))

    leaves: list[tuple[int, dict[str, object]]] = []
    stack: list[tuple[int, Decimal | None, int]] = []  # level, quantity per base, row
    product: str | None = None
    base = Decimal(1)
    assemblies = 0
    for position, (offset, level, values) in enumerate(entries):
        if level == 0:
            stack.clear()
            product = str(values["part_number"]) if values is not None else None
            quantity = values.get("qty_per_unit") if values is not None else None
            base = quantity if isinstance(quantity, Decimal) and quantity > 0 else Decimal(1)
            if values is not None:
                table.skipped_rows += 1
                table.total_rows -= 1
            continue
        while stack and stack[-1][0] >= level:
            stack.pop()
        above = stack[-1] if stack else None
        own = values.get("qty_per_unit") if values is not None else None
        per_parent = Decimal(1) if above is None else above[1]
        quantity = own * per_parent if isinstance(own, Decimal) and per_parent is not None else None
        stack.append((level, quantity, offset))
        is_assembly = position + 1 < len(entries) and entries[position + 1][1] > level
        if is_assembly:
            if values is not None:
                assemblies += 1
                table.skipped_rows += 1
                table.total_rows -= 1
            continue
        if values is None:
            continue  # already reported
        if quantity is None:
            parent = above[2] if above is not None else offset
            table.errors.append(RowError(offset, None, f"its assembly (row {parent}) is invalid"))
            continue
        leaf = dict(values)
        leaf["qty_per_unit"] = (quantity / base).normalize()
        if product is not None:
            leaf["product"] = product
        leaves.append((offset, leaf))
    table.rows = leaves
    levels = max((level for _, level, _ in entries), default=0)
    if levels < 2 and not assemblies:
        return []
    return [
        f"Multi-level BOM ({levels} levels) exploded: {len(leaves)} parts to buy, "
        f"{assemblies} assemblies not imported (built, not bought)."
    ]


async def load_workbook_upload(
    db: Db, storage: BlobStorage, principal: Principal, upload_id: UUID
) -> tuple[str, bytes]:
    """Read only an upload owned by the caller inside their tenant."""
    async with tenant_session(db, principal.tenant_id) as session:
        upload = await session.scalar(
            select(TemplateUpload).where(
                TemplateUpload.id == upload_id, TemplateUpload.user_id == principal.user_id
            )
        )
        if upload is None:
            raise WorkbookFileError(f"upload {upload_id} not found")
        key, filename = upload.storage_key, upload.filename
    return filename, await storage.get(key)


async def import_workbook(
    db: Db,
    storage: BlobStorage,
    principal: Principal,
    upload_id: UUID,
    *,
    mapper: SheetMapper | None,
    confirmed: bool = False,
    decimal: DecimalSeparator | None = None,
    mapping: SheetMappingOutput | None = None,
    dry_run: bool = False,
    apply_session: AsyncSession | None = None,
    source: tuple[str, bytes] | None = None,
) -> WorkbookImportReport:
    """Read every sheet, find the sections, and import them (or ask first)."""
    from bido_back.purchasing.service import apply_parsed

    filename, data = source or await load_workbook_upload(db, storage, principal, upload_id)
    try:
        grids = read_workbook(data, filename)
    except ImportFileError as exc:
        raise WorkbookFileError(str(exc)) from exc
    context = ToolContext(tenant_id=principal.tenant_id, user_id=principal.user_id, run_id=uuid4())
    by_name = {g.name: g for g in grids}
    if mapping is None:
        sections, problems = await map_workbook(grids, filename, mapper=mapper, context=context)
    else:
        sections, problems = _from_model(mapping, by_name)
    numbers = {id(s): _number_cells(s, by_name[s.sheet]) for s in sections}
    in_workbook = (
        decimal_evidence(c for cells in numbers.values() for c in cells)
        or _dotted_dates(sections, by_name)
        or decimal
    )
    parsed: list[_Parsed] = []
    for section in sections:
        try:
            table, notes = _parse_section(
                section,
                by_name[section.sheet],
                decimal_evidence(numbers[id(section)]) or in_workbook,
            )
        except ImportFileError as exc:
            problems.append(f"{section.sheet}: {exc}")
            continue
        parsed.append(_Parsed(section, table, [*notes, *_stock_notes(section, table)]))
    _names_for_codes(parsed)
    parsed_sheets = {p.section.sheet for p in parsed if p.table.rows}
    unimported = [
        g.name
        for g in grids
        if g.name not in parsed_sheets and any(not _blank(c) for row in g.rows for c in row)
    ]
    if not parsed:
        _log_import("nothing_found", upload_id, parsed, problems, sheets=len(grids))
        return WorkbookImportReport(
            upload_id=upload_id,
            filename=filename,
            status="nothing_found",
            sections=[],
            questions=["I could not find a table to import. Which sheet and columns hold it?"],
            problems=problems,
            unimported_sheets=unimported,
        )
    questions = [q for p in parsed for q in p.doubts()]
    if problems:
        questions.append("Some sections could not be read. Review the problems before importing.")
    if unimported:
        questions.append(
            f"These sheets have no importable rows: {', '.join(unimported)}. "
            "Import only the supported data?"
        )
    if dry_run or (questions and not confirmed):
        _log_import(
            "needs_confirmation",
            upload_id,
            parsed,
            problems,
            sheets=len(grids),
            questions=len(questions),
        )
        return WorkbookImportReport(
            upload_id=upload_id,
            filename=filename,
            status="needs_confirmation",
            sections=[p.report for p in parsed],
            questions=questions,
            problems=problems,
            unimported_sheets=unimported,
        )
    order = {t: i for i, t in enumerate(_ORDER)}

    async def write(session: AsyncSession) -> list[str]:
        for item in sorted(parsed, key=lambda p: order[p.section.import_type]):
            created, updated, skipped = await apply_parsed(session, principal.tenant_id, item.table)
            item.report.created, item.report.updated = created, updated
            item.report.skipped = skipped
        return await suppliers_without_email(session)

    if apply_session is None:
        async with tenant_session(db, principal.tenant_id) as session:
            without_email = await write(session)
    else:
        without_email = await write(apply_session)
    _log_import("imported", upload_id, parsed, problems, sheets=len(grids), confirmed=confirmed)
    return WorkbookImportReport(
        upload_id=upload_id,
        filename=filename,
        status="imported",
        sections=[p.report for p in parsed],
        questions=[],
        problems=problems,
        unimported_sheets=unimported,
        suppliers_without_email=without_email,
    )


async def suppliers_without_email(session: AsyncSession) -> list[str]:
    """The tenant's suppliers with no usable contact email, by name (bounded).

    A supplier master (SAP XK03) or a purchase history brings suppliers without
    contacts; an RFQ can only invite a supplier through a contact's email.
    """
    has_email = exists().where(
        SupplierContact.supplier_id == Supplier.id,
        SupplierContact.contact_id == Contact.id,
        Contact.email.is_not(None),
        Contact.pseudonymized_at.is_(None),
    )
    names = await session.scalars(
        select(Supplier.name)
        .where(~has_email)
        .order_by(Supplier.name)
        .limit(MAX_SUPPLIERS_WITHOUT_EMAIL)
    )
    return list(names)


_NUMBER: Final = re.compile(r"\d[\d.,]*")


def _log_import(
    status: str,
    upload_id: UUID,
    parsed: Sequence[_Parsed],
    problems: Sequence[str],
    *,
    sheets: int,
    questions: int = 0,
    confirmed: bool = False,
) -> None:
    """A line for the import and one per section: types, counts and error reasons.

    Sheet names, headers and cell values are the buyer's content, so sections
    are numbered, and row errors are reduced to the field and the masked
    message ("quantity: not a valid number: '…'"), counted.
    """
    logger.info(
        "workbook import",
        extra={
            "upload_id": str(upload_id),
            "status": status,
            "confirmed": confirmed,
            "sheets": sheets,
            "sections": len(parsed),
            "questions": questions,
            "problems": len(problems),
            "rows_ok": sum(p.report.rows_ok for p in parsed),
            "rows_failed": sum(p.report.rows_failed for p in parsed),
        },
    )
    for index, item in enumerate(parsed):
        report = item.report
        reasons: dict[str, int] = {}
        for error in item.table.errors:
            text = f"{error.field + ': ' if error.field else ''}{error.message}"
            # An unquoted cell value is a number; mask those too.
            key = _NUMBER.sub("<n>", safe_reason(text, 120))
            reasons[key] = reasons.get(key, 0) + 1
        top = sorted(reasons.items(), key=lambda kv: -kv[1])[:5]
        logger.log(
            logging.WARNING if report.rows_failed or report.missing_required else logging.INFO,
            "workbook section",
            extra={
                "upload_id": str(upload_id),
                "section": index,
                "type": report.type,
                "source": report.source,
                "header_row": report.header_row,
                "confidence": report.confidence,
                "rows_total": report.rows_total,
                "rows_ok": report.rows_ok,
                "rows_failed": report.rows_failed,
                "rows_skipped": report.rows_skipped,
                "missing_required": report.missing_required,
                "columns_mapped": len(report.columns),
                "has_question": report.question is not None,
                "error_reasons": [f"{count}x {reason}" for reason, count in top],
            },
        )


def _number_cells(section: Section, grid: SheetGrid) -> list[object]:
    columns = [
        i
        for spec in SPECS[section.import_type]
        if spec.kind in ("number", "int") and (i := section.index.get(spec.name)) is not None
    ]
    end = len(grid.rows) if section.last_row is None else section.last_row + 1
    return [row[i] for row in grid.rows[section.first_row : end] for i in columns if i < len(row)]


_DOTTED_DATE: Final = re.compile(r"\d{1,2}\.\d{1,2}\.\d{2}(?:\d{2})?")


def _dotted_dates(
    sections: Sequence[Section], grids: Mapping[str, SheetGrid]
) -> DecimalSeparator | None:
    """A last hint: day-first dotted dates (08.10.2026) come with decimal commas."""
    for section in sections:
        columns = [
            i
            for spec in SPECS[section.import_type]
            if spec.kind == "date" and (i := section.index.get(spec.name)) is not None
        ]
        for row in grids[section.sheet].rows[section.first_row :]:
            for i in columns:
                cell = row[i] if i < len(row) else None
                if isinstance(cell, str) and _DOTTED_DATE.fullmatch(cell.strip()):
                    return ","
    return None


# Stock that cannot go into a build: left out of on-hand stock.
_UNAVAILABLE_STOCK: Final = frozenset(
    {
        "em contr qualidade", "em controle de qualidade", "estoque em qualidade",
        "quality inspection", "in quality inspection", "qi stock", "estoque bloqueado",
        "bloqueado", "blocked", "blocked stock", "bloccato", "stock bloccato",
        "in controllo qualita", "controllo qualita",
    }
)  # fmt: skip


def _stock_notes(section: Section, table: ParsedTable) -> list[str]:
    if section.import_type != "inventory":
        return []
    left_out = [c for c in table.unknown_columns if norm(c) in _UNAVAILABLE_STOCK]
    if not left_out:
        return []
    return [
        f"Only unrestricted stock counted as on hand ({table.columns.get('on_hand')}); "
        f"left out: {', '.join(left_out)}."
    ]


def _names_for_codes(parsed: Sequence[_Parsed]) -> None:
    """Rows naming a supplier by its ERP code get the name another sheet gives it."""
    names: dict[str, str] = {}
    for item in parsed:
        for _, row in item.table.rows:
            code, name = row.get("supplier_code"), row.get("supplier")
            if isinstance(code, str) and isinstance(name, str) and code.lower() != name.lower():
                names.setdefault(code.lower(), name)
    if not names:
        return
    for item in parsed:
        renamed = 0
        for _, row in item.table.rows:
            name = row.get("supplier")
            if isinstance(name, str) and name.lower() in names:
                row["supplier"] = names[name.lower()]
                renamed += 1
        if renamed:
            item.notes.append(f"{renamed} rows named suppliers by code; matched to their names.")
            item.refresh()


def default_sheet_mapper(config: RunnerConfig | None = None) -> SheetMapper:
    async def run_mapper(task: SheetMappingInput, context: ToolContext) -> SheetMappingOutput:
        from bido_back.agents.catalog.definitions import SHEET_MAPPER
        from bido_back.agents.runner import Completed, Failed, run_to_completion

        result = await run_to_completion(SHEET_MAPPER, task, context, config=config)
        if isinstance(result, Completed):
            return result.output
        code = result.failure.code if isinstance(result, Failed) else "unexpected_approval"
        logger.error("agent sheet_mapper run %s failed: %s", context.run_id, code)
        return SheetMappingOutput(sections=[])

    return run_mapper
