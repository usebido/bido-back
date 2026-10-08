"""The fixtures: datasets written as the spreadsheets a buyer (or ChatGPT) produces.

Each fixture writes one dataset in one shape (tidy sheets, Italian headers with
title rows and a totals row, Portuguese with decimal commas, one sheet mixing
two tables, off-synonym headers, mixed date formats, duplicates and bad rows,
no BOM, no production plan, two products, a SAP export) and declares what a
correct import finds. The analysis golden comes from `evals.golden`.
"""

import io
from collections.abc import Callable, Sequence
from dataclasses import dataclass, field, replace
from datetime import date, timedelta
from decimal import Decimal
from typing import Any, Final, Literal, cast

from bido_back.agents.catalog.contracts import ColumnAssignment, SheetSection
from bido_back.purchasing.importers import ImportType
from evals.data import Dataset, PurchaseRow, build
from evals.erp import FILENAME, sap_motor
from evals.erp_wide import FILENAME as WIDE_FILENAME
from evals.erp_wide import wide_export

Language = Literal["en", "it", "pt"]

_EXCEL_EPOCH: Final = date(1899, 12, 30)


@dataclass(frozen=True, slots=True)
class Fixture:
    name: str
    description: str
    filename: str
    data: bytes
    dataset: Dataset
    language: Language
    # Units to analyse when the spreadsheet has no production plan.
    quantity: Decimal | None
    # (sheet, type) sections a correct import finds, and valid rows per type.
    sections: tuple[tuple[str, ImportType], ...]
    rows_ok: dict[ImportType, int]
    contacts: bool
    analysable: bool
    # Purchase-order lines a correct import stores (exact duplicates merged).
    unique_po_lines: int
    # The right mapping for sheets the rules cannot read (scripted mapper in mock mode).
    mapping: tuple[SheetSection, ...] = field(default_factory=tuple[SheetSection, ...])
    # The import must ask before writing (the buyer confirms in a later message).
    asks_confirmation: bool = False


# --- Writing -----------------------------------------------------------------------------


Sheet = tuple[str, list[list[object]], dict[str, Any]]


def _xlsx(sheets: Sequence[Sheet]) -> bytes:
    from openpyxl import Workbook  # pyright: ignore[reportMissingTypeStubs]

    workbook = cast(Any, Workbook())
    workbook.remove(workbook.active)
    for title, rows, options in sheets:
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
        for merged in options.get("merge", ()):
            sheet.merge_cells(merged)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _num(value: Decimal) -> float | int:
    return int(value) if value == value.to_integral_value() else float(value)


def _comma(value: Decimal, places: int = 2) -> str:
    """Brazilian format: 1.234,56 (thousands dot, decimal comma)."""
    text = f"{value:,.{places}f}"
    return text.replace(",", "§").replace(".", ",").replace("§", ".")


def _titled(title: str, header: list[object], rows: list[list[object]]) -> Sheet:
    width = len(header)
    last = chr(ord("A") + width - 1)
    return (
        title[:31],
        [[title], [], header, *rows],
        {"merge": [f"A1:{last}1"]},
    )


# --- Headers ----------------------------------------------------------------------------

_H: Final[dict[Language, dict[str, list[object]]]] = {
    "en": {
        "bom": ["Product", "Part Number", "Description", "MPN", "Manufacturer", "Unit",
                "Qty per unit"],
        "stock": ["Part Number", "Description", "On Hand", "On Order", "Location"],
        "po": ["PO Number", "Order Date", "Supplier", "Part Number", "Description", "Quantity",
               "Unit Price", "Currency", "Promised Date", "Received Date", "Rejected Qty"],
        "suppliers": ["Supplier", "Contact Name", "Email", "Country"],
        "plan": ["Product", "Quantity", "Need By"],
    },
    "it": {
        "bom": ["Prodotto", "Codice articolo", "Descrizione", "Codice produttore", "Produttore",
                "Unità", "Qtà per unità"],
        "stock": ["Codice articolo", "Descrizione", "Giacenza", "Ordinato", "Magazzino"],
        "po": ["Nr ordine", "Data ordine", "Fornitore", "Codice articolo", "Descrizione",
               "Quantità", "Prezzo unitario", "Valuta", "Data consegna prevista",
               "Data ricevimento", "Scarti", "Email fornitore"],
        "plan": ["Prodotto", "Quantità", "Data produzione"],
    },
    "pt": {
        "bom": ["Produto", "Código", "Descrição", "Código fabricante", "Fabricante", "Unidade",
                "Qtd por unidade"],
        "stock": ["Código", "Descrição", "Estoque", "Pedidos abertos", "Depósito"],
        "po": ["Nº Pedido", "Data do pedido", "Fornecedor", "Código", "Descrição", "Quantidade",
               "Preço unitário", "Moeda", "Data prevista", "Data entrega", "Qtd rejeitada"],
        "suppliers": ["Fornecedor", "Contato", "E-mail", "País"],
        "plan": ["Produto", "Quantidade", "Data necessidade"],
    },
}  # fmt: skip


def _bom_rows(ds: Dataset, numbers: Callable[[Decimal], object]) -> list[list[object]]:
    rows: list[list[object]] = [
        [ds.product, p.number, p.description, p.mpn or "", p.manufacturer or "", p.unit,
         numbers(p.qty_per_unit)]
        for p in ds.parts
    ]  # fmt: skip
    for product, bom in ds.other_boms.items():
        for number, qty in bom.items():
            p = ds.part(number)
            rows.append([product, p.number, p.description, p.mpn or "", p.manufacturer or "",
                         p.unit, numbers(qty)])  # fmt: skip
    return rows


def _stock_rows(ds: Dataset, numbers: Callable[[Decimal], object]) -> list[list[object]]:
    return [
        [number, ds.part(number).description, numbers(held[0]), numbers(held[1]), "Main"]
        for number, held in ds.stock.items()
    ]


def _po_row(
    row: PurchaseRow,
    numbers: Callable[[Decimal], object],
    prices: Callable[[Decimal], object],
    dates: Callable[[date, int], object],
    index: int,
) -> list[object]:
    return [
        row.po_number,
        dates(row.ordered_at, index),
        row.supplier,
        row.part,
        row.description,
        numbers(row.quantity),
        prices(row.unit_price),
        row.currency,
        dates(row.promised_at, index + 1),
        dates(row.received_at, index + 2) if row.received_at else "",
        numbers(row.rejected_qty) if row.rejected_qty else "",
    ]


def _supplier_rows(ds: Dataset) -> list[list[object]]:
    return [[s.name, s.contact, s.email, s.country] for s in ds.suppliers]


def _plan_rows(ds: Dataset, dates: Callable[[date, int], object]) -> list[list[object]]:
    return [[p.product, _num(p.quantity), dates(p.need_by, i)] for i, p in enumerate(ds.plan)]


def _as_date(value: date, _: int) -> object:
    return value


def _count(ds: Dataset) -> dict[ImportType, int]:
    bom = len(ds.parts) + sum(len(b) for b in ds.other_boms.values())
    return {
        "bom": bom,
        "inventory": len(ds.stock),
        "purchase_history": len(ds.purchases),
        "supplier_contacts": len(ds.suppliers),
        "production_plan": len(ds.plan),
    }


# --- The fixtures -------------------------------------------------------------------------


def _tidy(ds: Dataset, lang: Language = "en", *, dates: Callable[[date, int], object] = _as_date
          ) -> list[Sheet]:  # fmt: skip
    h = _H[lang]
    stock_header = h["stock"]
    stock_rows = _stock_rows(ds, _num)
    if ds.stock_as_of is not None:
        stock_header = [*stock_header, "As of", "On order due"]
        stock_rows = [[*row, ds.stock_as_of, ds.stock_due] for row in stock_rows]
    sheets: list[Sheet] = [
        ("BOM", [h["bom"], *_bom_rows(ds, _num)], {}),
        ("Stock", [stock_header, *stock_rows], {}),
        (
            "Purchases",
            [[*h["po"], "Received qty"]]
            + [
                [
                    *_po_row(r, _num, _num, dates, i),
                    _num(r.received_qty) if r.received_qty is not None else None,
                ]
                for i, r in enumerate(ds.purchases)
            ],
            {},
        ),  # fmt: skip
        ("Suppliers", [h["suppliers"], *_supplier_rows(ds)], {}),
    ]
    if ds.plan:
        sheets.append(("Plan", [h["plan"], *_plan_rows(ds, dates)], {}))
    return sheets


def _columns(*pairs: tuple[str, int]) -> list[ColumnAssignment]:
    return [ColumnAssignment(field=f, column=c) for f, c in pairs]


def _sections(*pairs: tuple[str, ImportType]) -> tuple[tuple[str, ImportType], ...]:
    return pairs


def fixtures(today: date) -> list[Fixture]:
    out: list[Fixture] = []

    # 1. The tidy baseline: one sheet per table, English headers on row 1.
    ds = build("motor_tidy_en", "motor", today, seed=1)
    ds = replace(
        ds,
        stock_as_of=today,
        stock_due=today + timedelta(days=15),
        purchases=tuple(
            replace(row, received_qty=row.quantity if row.received_at else None)
            for row in ds.purchases
        ),
    )
    out.append(
        Fixture(
            name=ds.key,
            description="One sheet per table, English headers on row 1, real numbers and dates",
            filename="motor_purchasing_data.xlsx",
            data=_xlsx(_tidy(ds)),
            dataset=ds,
            language="pt",
            quantity=None,
            sections=_sections(("BOM", "bom"), ("Stock", "inventory"),
                               ("Purchases", "purchase_history"),
                               ("Suppliers", "supplier_contacts"), ("Plan", "production_plan")),
            rows_ok=_count(ds),
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 2. Italian, a merged title row and a blank row above each table, a totals row.
    ds = build("battery_titles_it", "battery", today, seed=2)
    h = _H["it"]
    po_rows = [_po_row(r, _num, _num, _as_date, i) + [ds.supplier(r.supplier).email]
               for i, r in enumerate(ds.purchases)]  # fmt: skip
    total_qty = sum((r.quantity for r in ds.purchases), Decimal(0))
    po_rows.append(["TOTALE", "", "", "", "", _num(total_qty), "", "", "", "", "", ""])
    sheets = [
        _titled(f"Distinta base — {ds.product}", h["bom"], _bom_rows(ds, _num)),
        _titled("Giacenze di magazzino", h["stock"], _stock_rows(ds, _num)),
        _titled("Storico acquisti 12 mesi", h["po"], po_rows),
        _titled("Piano di produzione", h["plan"], _plan_rows(ds, _as_date)),
    ]
    named = [(f"Distinta base — {ds.product}"[:31], "bom"), ("Giacenze di magazzino", "inventory"),
             ("Storico acquisti 12 mesi", "purchase_history"),
             ("Storico acquisti 12 mesi", "supplier_contacts"),
             ("Piano di produzione", "production_plan")]  # fmt: skip
    rows_ok = _count(ds)
    rows_ok["supplier_contacts"] = len(ds.purchases)  # one contact row per purchase line
    out.append(
        Fixture(
            name=ds.key,
            description="Italian headers, merged title + blank row above each header, "
            "supplier e-mail next to each purchase, TOTALE row at the end",
            filename="pacco_batteria_acquisti.xlsx",
            data=_xlsx(sheets),
            dataset=ds,
            language="it",
            quantity=None,
            sections=tuple((s, cast(ImportType, t)) for s, t in named),
            rows_ok=rows_ok,
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )

    # 3. Portuguese with decimal commas, euro signs and dd/mm/yyyy text dates.
    ds = build("inverter_pt_comma", "inverter", today, seed=3)
    h = _H["pt"]

    def comma_qty(v: Decimal) -> object:
        return _comma(v, 0) if v == v.to_integral_value() else _comma(v, 2)

    def euro(v: Decimal) -> object:
        return f"€ {_comma(v, 2)}"

    def br_date(v: date, _: int) -> object:
        return v.strftime("%d/%m/%Y")

    sheets: list[Sheet] = [
        ("Lista de materiais", [h["bom"], *_bom_rows(ds, comma_qty)], {}),
        ("Estoque", [h["stock"], *_stock_rows(ds, comma_qty)], {}),
        (
            "Histórico de compras",
            [h["po"]]
            + [_po_row(r, comma_qty, euro, br_date, i) for i, r in enumerate(ds.purchases)],
            {},
        ),  # fmt: skip
        ("Fornecedores", [h["suppliers"], *_supplier_rows(ds)], {}),
        (
            "Plano de produção",
            [h["plan"]]
            + [[p.product, comma_qty(p.quantity), br_date(p.need_by, 0)] for p in ds.plan],
            {},
        ),  # fmt: skip
    ]
    out.append(
        Fixture(
            name=ds.key,
            description="Portuguese headers, numbers as text with thousands dot and decimal "
            "comma (1.200 / 36,50), '€ ' prefixes, dd/mm/yyyy text dates",
            filename="inversor_dados_compras.xlsx",
            data=_xlsx(sheets),
            dataset=ds,
            language="pt",
            quantity=None,
            sections=_sections(("Lista de materiais", "bom"), ("Estoque", "inventory"),
                               ("Histórico de compras", "purchase_history"),
                               ("Fornecedores", "supplier_contacts"),
                               ("Plano de produção", "production_plan")),
            rows_ok=_count(ds),
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 4. Two tables per sheet: purchases + supplier e-mail, BOM + stock; no plan.
    ds = build("motor_combined", "motor", today, seed=4, with_plan=False)
    h = _H["en"]
    po_header = [*h["po"], "Email"]
    po_rows = [_po_row(r, _num, _num, _as_date, i) + [ds.supplier(r.supplier).email]
               for i, r in enumerate(ds.purchases)]  # fmt: skip
    parts_header: list[object] = ["Product", "Part Number", "Description", "Unit", "Qty per unit",
                                  "On Hand", "On Order"]  # fmt: skip
    parts_rows: list[list[object]] = [
        [ds.product, p.number, p.description, p.unit, _num(p.qty_per_unit),
         _num(ds.stock[p.number][0]), _num(ds.stock[p.number][1])]
        for p in ds.parts
    ]  # fmt: skip
    rows_ok = _count(ds)
    rows_ok["supplier_contacts"] = len(ds.purchases)
    out.append(
        Fixture(
            name=ds.key,
            description="Two tables per sheet (purchases + supplier e-mail; BOM + stock), "
            "no production plan (quantity given)",
            filename="motor_all_in_two_sheets.xlsx",
            data=_xlsx([("Data", [po_header, *po_rows], {}),
                        ("Parts & stock", [parts_header, *parts_rows], {})]),
            dataset=ds,
            language="en",
            quantity=Decimal(150),
            sections=_sections(("Data", "purchase_history"), ("Data", "supplier_contacts"),
                               ("Parts & stock", "bom"), ("Parts & stock", "inventory")),
            rows_ok=rows_ok,
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 5. Only stock and purchases: nothing says what to build.
    ds = build("battery_no_bom", "battery", today, seed=5, with_plan=False)
    po_rows = [_po_row(r, _num, _num, _as_date, i) + [ds.supplier(r.supplier).email]
               for i, r in enumerate(ds.purchases)]  # fmt: skip
    rows_ok = _count(ds)
    rows_ok["supplier_contacts"] = len(ds.purchases)
    out.append(
        Fixture(
            name=ds.key,
            description="Only stock and purchase history (no BOM, no plan): the analysis "
            "cannot run; the copilot should ask what to build",
            filename="battery_stock_and_purchases.xlsx",
            data=_xlsx([("Stock", [h["stock"], *_stock_rows(ds, _num)], {}),
                        ("Purchases", [po_header, *po_rows], {})]),
            dataset=ds,
            language="pt",
            quantity=None,
            sections=_sections(("Stock", "inventory"), ("Purchases", "purchase_history"),
                               ("Purchases", "supplier_contacts")),
            rows_ok={k: v for k, v in rows_ok.items() if k != "bom"},
            contacts=True,
            analysable=False,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 6. Tidy English without a production plan: the quantity comes from the buyer.
    ds = build("inverter_no_plan", "inverter", today, seed=6, with_plan=False)
    out.append(
        Fixture(
            name=ds.key,
            description="Tidy English, no production plan (quantity 250 given by the buyer)",
            filename="inverter_data.xlsx",
            data=_xlsx(_tidy(ds)),
            dataset=ds,
            language="pt",
            quantity=Decimal(250),
            sections=_sections(("BOM", "bom"), ("Stock", "inventory"),
                               ("Purchases", "purchase_history"),
                               ("Suppliers", "supplier_contacts")),
            rows_ok={k: v for k, v in _count(ds).items() if k != "production_plan"},
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 7. Off-synonym headers under a note row: only the sheet mapper can read them.
    ds = build("motor_messy_headers", "motor", today, seed=7)
    messy: list[object] = ["PO #", "Date of Order", "Vendor Name", "Item Code",
                           "Item Description", "Qty Ordered", "Unit Cost (EUR)", "Promised By",
                           "Arrived On", "Defects", "Contact e-mail"]  # fmt: skip
    messy_rows: list[list[object]] = [
        [r.po_number, r.ordered_at, r.supplier, r.part, r.description, _num(r.quantity),
         _num(r.unit_price), r.promised_at, r.received_at or "",
         _num(r.rejected_qty) if r.rejected_qty else "", ds.supplier(r.supplier).email]
        for r in ds.purchases
    ]  # fmt: skip
    sheets: list[Sheet] = [
        ("BOM", [h["bom"], *_bom_rows(ds, _num)], {}),
        ("Stock", [h["stock"], *_stock_rows(ds, _num)], {}),
        (
            "ERP export",
            [[f"Exported from ERP on {today.isoformat()} — all amounts EUR"], messy, *messy_rows],
            {},
        ),
        ("Plan", [h["plan"], *_plan_rows(ds, _as_date)], {}),
    ]
    rows_ok = _count(ds)
    rows_ok["supplier_contacts"] = len(ds.purchases)
    mapping = (
        SheetSection(
            sheet="ERP export", import_type="purchase_history", header_row=2, first_data_row=3,
            last_data_row=None,
            columns=[ColumnAssignment(field=f, column=c) for f, c in (
                ("po_number", 0), ("ordered_at", 1), ("supplier", 2), ("part_number", 3),
                ("description", 4), ("quantity", 5), ("unit_price", 6), ("promised_at", 7),
                ("received_at", 8), ("rejected_qty", 9))],
            confidence=0.9, question=None,
        ),
        SheetSection(
            sheet="ERP export", import_type="supplier_contacts", header_row=2, first_data_row=3,
            last_data_row=None,
            columns=[ColumnAssignment(field="supplier", column=2),
                     ColumnAssignment(field="email", column=10)],
            confidence=0.9, question=None,
        ),
    )  # fmt: skip
    out.append(
        Fixture(
            name=ds.key,
            description="Purchases exported from an ERP with off-synonym headers (Vendor Name, "
            "Qty Ordered, Unit Cost (EUR), Date of Order) under a note row: needs the mapper",
            filename="motor_erp_export.xlsx",
            data=_xlsx(sheets),
            dataset=ds,
            language="en",
            quantity=None,
            sections=_sections(("BOM", "bom"), ("Stock", "inventory"),
                               ("ERP export", "purchase_history"),
                               ("ERP export", "supplier_contacts"), ("Plan", "production_plan")),
            rows_ok=rows_ok,
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
            mapping=mapping,
        )
    )  # fmt: skip

    # 8. Duplicated purchase lines and deliberately bad rows.
    ds = build("battery_dupes_bad_rows", "battery", today, seed=8)
    sheets = _tidy(ds)
    po_sheet = sheets[2][1]
    dupes = [list(row) for row in po_sheet[1:4]]
    bad: list[list[object]] = [
        ["PO-9001", today - timedelta(days=10), ds.suppliers[0].name, ds.parts[0].number,
         ds.parts[0].description, -50, 3.1, "EUR", "", "", ""],
        ["PO-9002", today - timedelta(days=11), ds.suppliers[0].name, ds.parts[0].number,
         ds.parts[0].description, "twenty", 3.1, "EUR", "", "", ""],
        ["PO-9003", today - timedelta(days=12), "", ds.parts[1].number,
         ds.parts[1].description, 10, 86, "EUR", "", "", ""],
    ]  # fmt: skip
    po_sheet.extend([*dupes, *bad])
    sheets[3][1].append(["Broken Contact Demo", "Mario Rossi", "rossi(at)example.test", "IT"])
    rows_ok = _count(ds)
    rows_ok["purchase_history"] = len(ds.purchases) + len(dupes)
    out.append(
        Fixture(
            name=ds.key,
            description="Tidy English plus 3 exact duplicate purchase lines, a negative "
            "quantity, a text quantity, a line without supplier and an invalid e-mail",
            filename="battery_data_with_errors.xlsx",
            data=_xlsx(sheets),
            dataset=ds,
            language="en",
            quantity=None,
            sections=_sections(("BOM", "bom"), ("Stock", "inventory"),
                               ("Purchases", "purchase_history"),
                               ("Suppliers", "supplier_contacts"), ("Plan", "production_plan")),
            rows_ok=rows_ok,
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 9. Mixed date formats in the same columns.
    ds = build("motor_mixed_dates", "motor", today, seed=9)

    def mixed(value: date, index: int) -> object:
        match index % 4:
            case 0:
                return value
            case 1:
                return value.isoformat()
            case 2:
                return value.strftime("%d.%m.%Y")
            case _:
                return (value - _EXCEL_EPOCH).days  # an Excel serial number

    out.append(
        Fixture(
            name=ds.key,
            description="Tidy English with dates as real dates, ISO text, dd.mm.yyyy text "
            "and Excel serial numbers mixed in the same columns",
            filename="motor_dates_mixed.xlsx",
            data=_xlsx(_tidy(ds, dates=mixed)),
            dataset=ds,
            language="en",
            quantity=None,
            sections=_sections(("BOM", "bom"), ("Stock", "inventory"),
                               ("Purchases", "purchase_history"),
                               ("Suppliers", "supplier_contacts"), ("Plan", "production_plan")),
            rows_ok=_count(ds),
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 10. Two products sharing parts in the BOM and the plan.
    ds = build("inverter_two_products", "inverter", today, seed=10, second_product=True)
    out.append(
        Fixture(
            name=ds.key,
            description="Two products sharing parts in one BOM and one plan; the analysis "
            "is for the main product only",
            filename="inverters_two_models.xlsx",
            data=_xlsx(_tidy(ds)),
            dataset=ds,
            language="en",
            quantity=None,
            sections=_sections(("BOM", "bom"), ("Stock", "inventory"),
                               ("Purchases", "purchase_history"),
                               ("Suppliers", "supplier_contacts"), ("Plan", "production_plan")),
            rows_ok=_count(ds),
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
        )
    )  # fmt: skip

    # 11-12. A SAP export (MB52, ME2M, CS12, MD04, XK03): titles above every table,
    # SAP field names, text numbers with three decimals, a three-level BOM named only
    # in its title, BRL and EUR, price units, vendor codes. The second also has stock
    # rows that cannot be read, so the import asks first (confirmed by text only).
    for gaps in (False, True):
        data, ds = sap_motor(today, gaps=gaps)
        rows_ok = _count(ds)
        out.append(
            Fixture(
                name=ds.key,
                description="SAP export: title rows, SAP field names, '888,000' quantities, "
                "3-level CS12 BOM with the product in the title, BRL+EUR, price unit 100, "
                "supplier master by vendor code"
                + (", plus unreadable stock rows (asks before importing)" if gaps else ""),
                filename=FILENAME,
                data=data,
                dataset=ds,
                language="pt",
                quantity=None,
                sections=_sections(("MB52 Estoque", "inventory"),
                                   ("ME2M Pedidos", "purchase_history"),
                                   ("CS12 Lista técnica", "bom"),
                                   ("Plano de produção", "production_plan"),
                                   ("XK03 Fornecedores", "supplier_contacts")),
                rows_ok=rows_ok,
                contacts=True,
                analysable=True,
                unique_po_lines=len(ds.purchases),
                asks_confirmation=gaps,
            )
        )  # fmt: skip

    # 13. A wide AI-generated "ERP DEMO" export: 46 English snake-case columns in the
    # purchase history, a UoM column (EA/KG) that is not a number, a local price next
    # to a USD one, cancelled lines, open lines with an outstanding quantity, a notes
    # sheet; stock, BOM and MRP plan sheets with off-synonym headers (the mapper).
    data, ds = wide_export(today)
    mapping = (
        SheetSection(sheet="Estoque_MMBE", import_type="inventory", header_row=1,
                     first_data_row=2, last_data_row=None,
                     columns=_columns(("part_number", 0), ("location", 3), ("on_hand", 4)),
                     confidence=0.9, question=None),
        SheetSection(sheet="Estrutura_BOM", import_type="bom", header_row=1, first_data_row=2,
                     last_data_row=None,
                     columns=_columns(("product", 0), ("part_number", 1), ("description", 2),
                                     ("qty_per_unit", 3), ("unit", 4), ("mpn", 5),
                                     ("manufacturer", 6)),
                     confidence=0.9, question=None),
        SheetSection(sheet="MRP_Plano", import_type="production_plan", header_row=1,
                     first_data_row=2, last_data_row=None,
                     columns=_columns(("product", 0), ("quantity", 1), ("need_by", 2)),
                     confidence=0.9, question=None),
        SheetSection(sheet="LEIA-ME", import_type="ignore", header_row=1, first_data_row=2,
                     last_data_row=None, columns=[], confidence=0.9, question=None),
    )  # fmt: skip
    out.append(
        Fixture(
            name=ds.key,
            description="Wide AI-generated ERP export: 46-column purchase history with a UoM "
            "column, local and USD prices, cancelled and open lines; stock, BOM and MRP plan "
            "with off-synonym headers; a notes sheet",
            filename=WIDE_FILENAME,
            data=data,
            dataset=ds,
            language="pt",
            quantity=None,
            sections=_sections(("Historico_Compras", "purchase_history"),
                               ("Estoque_MMBE", "inventory"), ("Estrutura_BOM", "bom"),
                               ("MRP_Plano", "production_plan"),
                               ("Fornecedores", "supplier_contacts")),
            rows_ok=_count(ds),
            contacts=True,
            analysable=True,
            unique_po_lines=len(ds.purchases),
            mapping=mapping,
        )
    )  # fmt: skip
    return out
