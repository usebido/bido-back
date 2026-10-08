"""A wide ERP export of the kind people generate with an AI assistant ("ERP DEMO").

Shaped after a real buyer test (2026-10-08): a purchase-history sheet with 40+
English snake-case columns, among them a unit-of-measure column (EA, KG) that
must never be read as a number, a local price next to a converted USD price,
cancelled lines (status CANCELLED, cancelled quantity = ordered), open lines
with an outstanding quantity, invoice and payment columns, and a notes sheet
("LEIA-ME") to ignore; plus stock, BOM, MRP plan and supplier sheets in the
same style. Supplier e-mails use reserved test domains only.

`wide_export()` writes the workbook and returns the dataset it means: the
cancelled lines are not in it (they are not purchases), and stock has nothing
on order (open purchase lines carry it). Run as a script to write the workbook
for a manual test:

    uv run python -m evals.erp_wide ~/Downloads/bido-teste-erp-wide.xlsx
"""

import io
import random
import sys
from dataclasses import replace
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Final, cast

from evals.data import Dataset, PurchaseRow, build

FILENAME: Final = "ERP_DEMO_Compras_Estoque.xlsx"
KEY: Final = "erp_wide_export"
_USD_PER_EUR: Final = Decimal("1.10")
_CENT: Final = Decimal("0.01")
_UOM: Final = {"pcs": "EA", "kg": "KG", "m": "M"}

PO_HEADER: Final[list[object]] = [
    "Company_Code", "PO_Number", "PO_Line", "PO_Date", "Supplier_ID", "Supplier_Name",
    "Material_ID", "Material_Description", "Plant", "Buyer_Group", "Currency", "FX_to_USD",
    "Ordered_Qty", "UoM", "Unit_Price_Local", "Unit_Price_USD", "Standard_Cost_USD",
    "Price_Variance_Pct", "Ordered_Value_USD", "Cancelled_Qty", "Net_Order_Value_USD",
    "Requested_Delivery_Date", "Confirmed_Delivery_Date", "Last_Receipt_Date",
    "Received_Qty_Gross", "Rejected_Qty", "Net_Received_Qty", "Outstanding_Qty",
    "Receipt_Value_USD", "Open_Value_USD", "PO_Status", "Delivery_Delay_Days", "Due_Line_Flag",
    "OTIF_Flag", "Incoterm", "Payment_Terms_Days", "Invoice_Number", "Invoice_Date",
    "Invoiced_Qty", "Invoice_Value_USD", "Invoice_Due_Date", "Payment_Date", "Payment_Status",
    "GR_Document", "Quality_Issue_Code", "Record_Type",
]  # fmt: skip


def _usd(value: Decimal) -> Decimal:
    return (value * _USD_PER_EUR).quantize(_CENT, rounding=ROUND_HALF_UP)


def _po_line(
    row: PurchaseRow, line: int, codes: dict[str, str], uom: str, today: date, status: str
) -> list[object]:
    cancelled = row.quantity if status == "CANCELLED" else Decimal(0)
    gross = Decimal(0)
    if status != "CANCELLED" and row.received_at is not None:
        gross = row.received_qty if row.received_qty is not None else row.quantity
    net = gross - row.rejected_qty
    outstanding = row.quantity - cancelled - gross if status != "CANCELLED" else Decimal(0)
    usd = _usd(row.unit_price)
    delay = (row.received_at - row.promised_at).days if row.received_at else 0
    invoiced = row.received_at is not None and status != "CANCELLED"
    return [
        "DEMO", row.po_number, line, row.ordered_at, codes[row.supplier], row.supplier,
        row.part, row.description, "PL01", "B10", row.currency, float(_USD_PER_EUR),
        float(row.quantity), uom, float(row.unit_price), float(usd),
        float((usd * Decimal("1.02")).quantize(_CENT)), -0.02,
        float(usd * row.quantity), float(cancelled),
        float(usd * (row.quantity - cancelled)),
        row.promised_at - timedelta(days=5), row.promised_at, row.received_at,
        float(gross), float(row.rejected_qty), float(net), float(outstanding),
        float(usd * net), float(usd * outstanding), status, max(delay, 0), int(outstanding > 0),
        int(invoiced and delay <= 0), "FCA", 45,
        f"IV-{row.po_number}" if invoiced else None,
        row.received_at + timedelta(days=2) if invoiced and row.received_at else None,
        float(net) if invoiced else 0, float(usd * net) if invoiced else 0,
        row.received_at + timedelta(days=47) if invoiced and row.received_at else None,
        None, "OPEN" if invoiced else "NOT_INVOICED",
        f"GR-{row.po_number}" if gross else None, "NONE", "SYNTHETIC",
    ]  # fmt: skip


def wide_export(today: date, *, seed: int = 13) -> tuple[bytes, Dataset]:
    """The workbook and the dataset it means (cancelled lines left out)."""
    rng = random.Random(seed)  # noqa: S311 - synthetic data, not security
    ds = build(KEY, "battery", today, seed=seed)
    # Open purchase lines carry what is on order; the stock sheet has no such column.
    ds = replace(ds, stock={number: (held[0], Decimal(0)) for number, held in ds.stock.items()})
    codes = {s.name: f"SUP-{100001 + i}" for i, s in enumerate(ds.suppliers)}
    units = {p.number: _UOM.get(p.unit, "EA") for p in ds.parts}

    po_rows: list[list[object]] = []
    for i, row in enumerate(ds.purchases):
        status = "CLOSED" if row.received_at is not None else "OPEN"
        po_rows.append(_po_line(row, 10, codes, units[row.part], today, status))
        if i % 6 == 2:  # a cancelled second line on the same PO: not a purchase
            ghost = replace(row, received_at=None, rejected_qty=Decimal(0))
            po_rows.append(_po_line(ghost, 20, codes, units[row.part], today, "CANCELLED"))
    rng.shuffle(po_rows)

    stock_rows: list[list[object]] = [
        [number, ds.part(number).description, "PL01", "WH-A", float(held[0]), units[number],
         float(_usd(ds.part(number).base_price) * held[0])]
        for number, held in ds.stock.items()
    ]  # fmt: skip
    bom_rows: list[list[object]] = [
        [ds.product, p.number, p.description, float(p.qty_per_unit), units[p.number],
         p.mpn or "", p.manufacturer or ""]
        for p in ds.parts
    ]  # fmt: skip
    plan_rows: list[list[object]] = [
        [p.product, float(p.quantity), p.need_by, "PLANNED_ORDER", "PL01"] for p in ds.plan
    ]
    supplier_rows: list[list[object]] = [
        [codes[s.name], s.name, s.country, s.contact, s.email, 45, "ACTIVE"] for s in ds.suppliers
    ]
    notes: list[list[object]] = [
        ["ERP DEMO | COMPRAS E ESTOQUE", "EXPORT SIMULADO"],
        ["AVISO", "Dados 100% fictícios, gerados para teste."],
        ["Registros", len(po_rows)],
        ["Data de corte", today],
    ]
    sheets: list[tuple[str, list[object], list[list[object]]]] = [
        ("Historico_Compras", PO_HEADER, po_rows),
        ("Estoque_MMBE", ["Material_ID", "Material_Description", "Plant", "Storage_Location",
                          "Unrestricted_Qty", "UoM", "Stock_Value_USD"], stock_rows),
        ("Estrutura_BOM", ["Parent_Material", "Component_ID", "Component_Description",
                           "Qty_Per_Parent", "UoM", "Manufacturer_Part_No", "Manufacturer"],
         bom_rows),
        ("MRP_Plano", ["Material_ID", "Planned_Qty", "Requirement_Date", "MRP_Element",
                       "Plant"], plan_rows),
        ("Fornecedores", ["Supplier_ID", "Supplier_Name", "Country", "Contact_Name",
                          "Contact_Email", "Payment_Terms_Days", "Status"], supplier_rows),
    ]  # fmt: skip

    from openpyxl import Workbook  # pyright: ignore[reportMissingTypeStubs]

    workbook = cast(Any, Workbook())
    workbook.remove(workbook.active)
    for title, header, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    readme = workbook.create_sheet("LEIA-ME")
    for row in notes:
        readme.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), ds


def cancelled_lines(today: date, *, seed: int = 13) -> int:
    """Cancelled purchase lines the workbook adds (skipped by a correct import)."""
    ds = build(KEY, "battery", today, seed=seed)
    return sum(1 for i in range(len(ds.purchases)) if i % 6 == 2)


if __name__ == "__main__":
    out = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else Path(FILENAME)
    data, built = wide_export(date.today())
    out.write_bytes(data)
    print(f"{out}: {len(built.parts)} parts, {len(built.purchases)} PO lines")
