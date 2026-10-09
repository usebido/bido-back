"""The buyer journey's ERP export: the golden answer, and a small equivalent workbook.

The journey (`evals.journey`) starts from an ERP export of a motor plant: a
purchase-order register with closed and open orders, two stock snapshots (the
export day's and an older one that must not be counted), production orders and
requisitions of three areas, item, supplier and approved-source masters,
receipts with a pending-delivery block, supplier contacts, the BOM and the
finished-product master, and archived quotation e-mails.

Two workbooks share one golden answer (Motor MTR-100, Planta A, from the export
date to 14 days later, no extra safety stock): 000101=50, 000102=30 (the 20
pieces on order arrive too late), 000103=0 (covered), 000104=22 kg, 000105=220,
000106=60, bought before from MetalNova, PrecisaoSul and MultiPecas.

* `real_case(path)`: the synthetic export the diagnosis used
  (`Exportacao_ERP_Industrial.xlsx`, 2026-10-08). It is never committed; pass
  it with `--workbook`. Its reference prices below were computed from its
  closed orders with the legacy rule (target = min(last, quantity-weighted
  average), maximum = max(last, average) * 1.05).
* `mock_case(export)`: a smaller workbook generated here with the same stock,
  plans, BOM, open orders and last purchases (the two most recent closed
  orders per part and supplier), every date shifted to `export`. Its sheet
  mapping is known, so mock mode answers the source mapper by script.

Every number here is decided by code. Domains are reserved test domains only.
"""

import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, Final, cast

from bido_back.agents.context import ToolContext
from bido_back.purchasing.sources.contracts import (
    ColumnMapping,
    MappingProposal,
    MappingTask,
    RecordKind,
    TableProposal,
)
from bido_back.purchasing.sources.mapping import Mapper

XLSX: Final = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
REAL_FILENAME: Final = "Exportacao_ERP_Industrial.xlsx"
MOCK_FILENAME: Final = "Exportacao_ERP_Motores_eval.xlsx"
REAL_EXPORT: Final = date(2026, 10, 8)

PRODUCT: Final = "MTR-100"
PLANT: Final = "Planta A"
HORIZON_DAYS: Final = 14
# The motor build in the horizon (OP-260884) is needed 10 days after the export.
NEED_DAYS: Final = 10
MAX_OVER: Final = Decimal("1.05")
CENT: Final = Decimal("0.01")


# --- Who sells what ---------------------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class SupplierRef:
    code: str
    name: str  # as the workbook writes it
    key: str  # how an answer may write it, letters only and lower case
    email: str

    @property
    def domain(self) -> str:
        return self.email.rpartition("@")[2]


METALNOVA: Final = SupplierRef(
    "F001", "MetalNova Teste", "metalnova", "vendas@metalnova.example.test"
)
PRECISAOSUL: Final = SupplierRef(
    "F002", "PrecisaoSul Teste", "precisaosul", "vendas@precisaosul.example.test"
)
MULTIPECAS: Final = SupplierRef(
    "F003", "MultiPecas Teste", "multipecas", "vendas@multipecas.example.test"
)
SUPPLIERS: Final = (METALNOVA, PRECISAOSUL, MULTIPECAS)


@dataclass(frozen=True, slots=True)
class PartGolden:
    """One component of the motor: what to buy, and what the history says about it."""

    part: str
    description: str
    unit: str
    to_buy: Decimal
    # Supplier key -> unit price of its last RECEIVED (closed) order.
    last_prices: Mapping[str, Decimal]
    # An open order's price: never the last price paid.
    open_price: Decimal | None
    # Internal reference (legacy rule), used when the product sets no limits.
    target: Decimal
    maximum: Decimal

    @property
    def buy(self) -> bool:
        return self.to_buy > 0


def _d(value: str) -> Decimal:
    return Decimal(value)


REAL_PARTS: Final = (
    PartGolden("000101", "Rolamento", "pcs", _d("50"),
               {"metalnova": _d("40.59"), "precisaosul": _d("49.44")}, _d("49.05"),
               _d("45.67"), _d("51.91")),
    PartGolden("000102", "Eixo", "pcs", _d("30"),
               {"precisaosul": _d("263.48"), "multipecas": _d("234.52")}, _d("240.00"),
               _d("234.52"), _d("266.28")),
    PartGolden("000103", "Carcaça", "pcs", _d("0"),
               {"metalnova": _d("166.71"), "multipecas": _d("184.70")}, None,
               _d("177.31"), _d("193.94")),
    PartGolden("000104", "Fio de cobre", "kg", _d("22"),
               {"metalnova": _d("109.37"), "precisaosul": _d("128.86")}, None,
               _d("116.27"), _d("135.30")),
    PartGolden("000105", "Ímã", "pcs", _d("220"),
               {"precisaosul": _d("23.37"), "multipecas": _d("20.81")}, _d("22.89"),
               _d("20.81"), _d("23.51")),
    PartGolden("000106", "Parafuso", "pcs", _d("60"),
               {"metalnova": _d("2.37"), "multipecas": _d("2.57")}, None,
               _d("2.49"), _d("2.70")),
)  # fmt: skip


@dataclass(frozen=True, slots=True)
class Case:
    """A workbook and its golden answer."""

    name: str
    filename: str
    data: bytes
    export: date
    parts: tuple[PartGolden, ...]
    # The mapping a correct source mapper proposes (mock); None: the real mapper runs.
    mapper: Mapper | None = None

    @property
    def horizon_end(self) -> date:
        return self.export + timedelta(days=HORIZON_DAYS)

    @property
    def need_by(self) -> date:
        return self.export + timedelta(days=NEED_DAYS)

    @property
    def to_buy(self) -> tuple[PartGolden, ...]:
        return tuple(p for p in self.parts if p.buy)

    def part(self, code: str) -> PartGolden:
        return next(p for p in self.parts if p.part == code)

    def scope(self) -> dict[str, object]:
        """The planning scope the buyer's message states, as tool arguments."""
        return {
            "product_code": PRODUCT,
            "location": PLANT,
            "as_of": self.export.isoformat(),
            "horizon_end": self.horizon_end.isoformat(),
            "safety_stock": "0",
        }


def real_case(path: Path) -> Case:
    """The synthetic export the diagnosis used (never committed)."""
    return Case(
        name="erp_export_mtr100",
        filename=path.name,
        data=path.read_bytes(),
        export=REAL_EXPORT,
        parts=REAL_PARTS,
    )


# --- Reference prices (the legacy rule, written from its docstring) ---------------------------


@dataclass(frozen=True, slots=True)
class PastOrder:
    po: str
    part: str
    supplier: SupplierRef
    quantity: Decimal
    price: Decimal
    ordered: date
    promised: date
    received: date | None = None
    received_qty: Decimal | None = None
    rejected: Decimal = Decimal(0)
    plant: str = PLANT


def reference_prices(orders: Sequence[PastOrder]) -> tuple[Decimal, Decimal]:
    """(target, maximum): target = min(last, average), maximum = max(last, average) * 1.05.

    Over the received orders of one part; the average is quantity-weighted and the
    last is the most recently ordered. Both are rounded to cents.
    """
    received = sorted(
        (o for o in orders if o.received is not None), key=lambda o: (o.ordered, o.po)
    )
    if not received:
        raise ValueError("no received order")
    quantity = sum((o.quantity for o in received), Decimal(0))
    average = sum((o.quantity * o.price for o in received), Decimal(0)) / quantity
    last = received[-1].price
    target = min(last, average).quantize(CENT, rounding=ROUND_HALF_UP)
    maximum = (max(last, average) * MAX_OVER).quantize(CENT, rounding=ROUND_HALF_UP)
    return target, maximum


# --- The equivalent workbook ------------------------------------------------------------------

_BASE: Final = REAL_EXPORT


def _on(day: str) -> date:
    """A date of the original export (YYYY-MM-DD), as an offset kept by `_Shift`."""
    return date.fromisoformat(day)


@dataclass(frozen=True, slots=True)
class _Shift:
    export: date

    def __call__(self, day: date) -> date:
        return self.export + (day - _BASE)

    def cell(self, day: date) -> datetime:
        return datetime.combine(self(day), time())


MN, PS, MP = METALNOVA, PRECISAOSUL, MULTIPECAS

# The two most recent closed orders per part and supplier, as in the original export:
# order, part, supplier, quantity, unit price, ordered, promised, received, received
# quantity, rejected quantity.
_CLOSED_TABLE: Final = """
510009 000101 F001 65   41.53  2026-08-03 2026-08-13 2026-08-16 65   7
510011 000101 F001 15   40.59  2026-08-28 2026-09-08 2026-09-13 15   1
510010 000101 F002 60   48.12  2026-08-14 2026-08-24 2026-08-22 60   0
510012 000101 F002 55   49.44  2026-09-12 2026-09-21 2026-09-22 55   0
510021 000102 F002 29   264.87 2026-08-02 2026-08-10 2026-08-10 29   0
510023 000102 F002 49   263.48 2026-08-29 2026-09-05 2026-09-04 49   0
510022 000102 F003 25   239.05 2026-08-14 2026-08-22 2026-08-21 25   2
510024 000102 F003 26   234.52 2026-09-13 2026-09-25 2026-09-26 26   0
510033 000103 F001 49   176.40 2026-08-04 2026-08-12 2026-08-15 49   4
510035 000103 F001 54   166.71 2026-09-01 2026-09-13 2026-09-13 54   4
510034 000103 F003 48   183.46 2026-08-18 2026-08-26 2026-08-27 48   5
510036 000103 F003 41   184.70 2026-09-10 2026-09-21 2026-09-23 41   0
510045 000104 F001 19.4 107.76 2026-08-04 2026-08-12 2026-08-12 19.4 1.5
510047 000104 F001 27.3 109.37 2026-08-28 2026-09-06 2026-09-07 27.3 1.4
510046 000104 F002 38.6 128.98 2026-08-13 2026-08-20 2026-08-20 38.6 0
510048 000104 F002 48.8 128.86 2026-09-09 2026-09-15 2026-09-15 48.8 0
510057 000105 F002 170  23.03  2026-07-31 2026-08-12 2026-08-10 170  0
510059 000105 F002 140  23.37  2026-08-30 2026-09-09 2026-09-09 140  0
510058 000105 F003 170  21.42  2026-08-15 2026-08-22 2026-08-22 170  12
510060 000105 F003 60   20.81  2026-09-12 2026-09-19 2026-09-21 60   0
510069 000106 F001 320  2.42   2026-08-05 2026-08-16 2026-08-16 320  38
510071 000106 F001 280  2.37   2026-08-31 2026-09-13 2026-09-14 280  9
510070 000106 F003 200  2.56   2026-08-15 2026-08-28 2026-08-28 200  6
510072 000106 F003 250  2.57   2026-09-13 2026-09-20 2026-09-22 250  0
"""


def _orders(table: str) -> tuple[PastOrder, ...]:
    by_code = {s.code: s for s in SUPPLIERS}
    orders: list[PastOrder] = []
    for row in table.strip().splitlines():
        po, part, code, qty, price, ordered, promised, received, got, rejected = row.split()
        orders.append(
            PastOrder(
                po=po,
                part=part,
                supplier=by_code[code],
                quantity=Decimal(qty),
                price=Decimal(price),
                ordered=_on(ordered),
                promised=_on(promised),
                received=_on(received),
                received_qty=Decimal(got),
                rejected=Decimal(rejected),
            )
        )
    return tuple(orders)


_CLOSED: Final = _orders(_CLOSED_TABLE)
# Other areas' history (maintenance in Planta A, pumps in Planta B): never motor demand.
_OTHER_CLOSED: Final = (
    PastOrder("510093", "000302", MN, Decimal(37), Decimal("75.40"), _on("2026-07-03"),
              _on("2026-07-11"), _on("2026-07-11"), Decimal(37)),
    PastOrder("510079", "000202", MP, Decimal(10), Decimal("18.64"), _on("2026-06-25"),
              _on("2026-07-07"), _on("2026-07-07"), Decimal(10), plant="Planta B"),
)  # fmt: skip
# Open orders: (po, part, supplier, quantity, price, ordered, promised = expected delivery).
_OPEN: Final = (
    ("520101", "000101", PS, "10", "49.05", "2026-10-05", "2026-10-13"),
    ("520102", "000102", MP, "20", "240.00", "2026-10-06", "2026-10-26"),
    ("520103", "000105", PS, "30", "22.89", "2026-10-07", "2026-10-13"),
)
_ITEMS: Final = (
    ("000101", "Rolamento", "pcs", "Motores"),
    ("000102", "Eixo", "pcs", "Motores"),
    ("000103", "Carcaça", "pcs", "Motores"),
    ("000104", "Fio de cobre", "kg", "Motores"),
    ("000105", "Ímã", "pcs", "Motores"),
    ("000106", "Parafuso", "pcs", "Motores"),
    ("000201", "Rotor bomba BH-40", "pcs", "Bombas hidráulicas"),
    ("000202", "Vedação bomba BH-40", "pcs", "Bombas hidráulicas"),
    ("000203", "Impulsor BH-40", "pcs", "Bombas hidráulicas"),
    ("000301", "Filtro linha ar", "pcs", "Manutenção"),
    ("000302", "Correia equipamento", "pcs", "Manutenção"),
    ("000303", "Óleo industrial ISO 46", "L", "Manutenção"),
)
# (plant, part, unit, gross, reserved, blocked) on the export day, then 19 days before.
_STOCK: Final = (
    ("Planta A", "000101", "pcs", 25, 5, 0),
    ("Planta A", "000102", "pcs", 15, 5, 0),
    ("Planta A", "000103", "pcs", 50, 0, 5),
    ("Planta A", "000104", "kg", 12, 2, 0),
    ("Planta A", "000105", "pcs", 100, 20, 10),
    ("Planta A", "000106", "pcs", 100, 0, 0),
    ("Planta A", "000302", "pcs", 22, 2, 1),
    ("Planta A", "000303", "L", 125, 15, 5),
    ("Planta B", "000202", "pcs", 155, 25, 0),
)
_OLD_STOCK: Final = (
    ("Planta A", "000101", "pcs", 35, 2, 0),
    ("Planta A", "000102", "pcs", 27, 5, 0),
    ("Planta A", "000103", "pcs", 49, 0, 5),
    ("Planta A", "000104", "kg", 23, 4, 0),
    ("Planta A", "000105", "pcs", 110, 22, 9),
    ("Planta A", "000106", "pcs", 106, 4, 0),
    ("Planta B", "000202", "pcs", 160, 28, 1),
)
# (order, plant, area, product, quantity, unit, needed, status, origin)
_PLANS: Final = (
    ("OP-260911", "Planta A", "Motores", "MTR-100", 30, "pcs", "2026-11-01", "Planejada",
     "Ordem fabricação"),
    ("OP-260875", "Planta A", "Manutenção", "MAN-KIT02", 8, "kits", "2026-11-07", "Planejada",
     "Ordem interna"),
    ("OP-260891", "Planta B", "Bombas hidráulicas", "BOM-BH40", 18, "pcs", "2026-10-17",
     "Planejada", "Ordem fabricação"),
    ("OP-260884", "Planta A", "Motores", "MTR-100", 40, "pcs", "2026-10-18", "Planejada",
     "Ordem fabricação"),
    ("OP-260919", "Planta A", "Manutenção", "MAN-KIT02", 12, "kits", "2026-10-20", "Liberada",
     "Ordem interna"),
)  # fmt: skip
_REQUISITIONS: Final = (
    ("RI-40081", "Planta A", "Manutenção", "000303", 35, "L", "2026-10-16"),
    ("RI-40086", "Planta B", "Bombas hidráulicas", "000202", 20, "pcs", "2026-10-21"),
)
_BOM: Final = (
    ("MAN-KIT02", "000302", 1, "pcs", "C", "2026-08-09"),
    ("BOM-BH40", "000202", 3, "pcs", "B", "2026-07-20"),
    ("MTR-100", "000103", 1, "pcs", "A", "2026-07-05"),
    ("MTR-100", "000105", 8, "pcs", "A", "2026-07-05"),
    ("MAN-KIT02", "000301", 2, "pcs", "C", "2026-08-09"),
    ("MTR-100", "000104", 0.8, "kg", "A", "2026-07-05"),
    ("MTR-100", "000102", 1, "pcs", "A", "2026-07-05"),
    ("MTR-100", "000106", 4, "pcs", "A", "2026-07-05"),
    ("MTR-100", "000101", 2, "pcs", "A", "2026-07-05"),
)
_PRODUCTS: Final = (
    ("MTR-100", "Motor MTR-100", "pcs", "Motores"),
    ("BOM-BH40", "Bomba hidráulica BH-40", "pcs", "Bombas hidráulicas"),
    ("MAN-KIT02", "Kit manutenção rotina", "kits", "Manutenção"),
)
# (part, supplier, approved since)
_APPROVED: Final = (
    ("000101", PS, "2025-11-30"), ("000101", MN, "2025-12-29"), ("000102", MP, "2025-12-11"),
    ("000102", PS, "2026-03-17"), ("000103", MP, "2026-03-09"), ("000103", MN, "2026-02-12"),
    ("000104", PS, "2026-03-02"), ("000104", MN, "2026-04-15"), ("000105", PS, "2026-01-22"),
    ("000105", MP, "2025-11-22"), ("000106", MP, "2026-05-25"), ("000106", MN, "2025-11-22"),
    ("000302", MN, "2025-11-29"), ("000202", MP, "2025-11-14"),
)  # fmt: skip
_CONTACTS: Final = (
    (MP, "Vendas internas", "Comercial", "08:30-17:30"),
    (MN, "Equipe comercial", "Comercial", "08:00-17:00"),
    (PS, "Atendimento B2B", "Comercial", "09:00-18:00"),
)


@dataclass(slots=True)
class _Block:
    """One table of a sheet: a title (except the first), a header, rows, and its mapping."""

    title: str | None
    header: tuple[str, ...]
    rows: list[tuple[object, ...]]
    kind: RecordKind | None = None
    # 1-based column -> canonical field (the mapping a correct source mapper proposes).
    fields: dict[int, str] = field(default_factory=dict[int, str])


@dataclass(slots=True)
class _Sheet:
    name: str
    title: str
    selection: str
    blocks: list[_Block]


def _sheets(shift: _Shift) -> list[_Sheet]:
    when = shift.cell
    pos: list[tuple[object, ...]] = []
    for o in (*_CLOSED, *_OTHER_CLOSED):
        pos.append((o.po, "00010", o.plant, o.part, o.supplier.code, float(o.quantity),
                    "kg" if o.part == "000104" else "pcs", float(o.price), "BRL",
                    when(o.ordered), when(o.promised), "Encerrado", "NB"))  # fmt: skip
    for po, part, supplier, quantity, price, ordered, promised in _OPEN:
        pos.append((po, "00010", PLANT, part, supplier.code, float(quantity), "pcs", float(price),
                    "BRL", when(_on(ordered)), when(_on(promised)), "Em aberto", "NB"))  # fmt: skip
    receipts: list[tuple[object, ...]] = [
        (o.po, "00010", o.plant, when(o.received), float(o.received_qty or 0), float(o.rejected),
         "kg" if o.part == "000104" else "pcs", "101", f"NF1{o.po[-4:]}")
        for o in (*_CLOSED, *_OTHER_CLOSED)
        if o.received is not None
    ]  # fmt: skip
    pending: list[tuple[object, ...]] = [
        (po, "00010", PLANT, part, float(quantity), "pcs", when(_on(promised)), "Não recebido",
         "Conf. fornecedor")
        for po, part, _supplier, quantity, _price, _ordered, promised in reversed(_OPEN)
    ]  # fmt: skip
    export = when(_BASE)
    old = when(_BASE - timedelta(days=19))
    return [
        _Sheet("ZMM_042", "MM / posição de pedidos de compra",
               "Centros A/B | classes NB | histórico e pendentes", [
            _Block(None, ("Pedido", "It.", "Centro destino", "Cód. material", "Forn.",
                          "Qtd. pedida", "UM", "Preço unit./UM", "Moeda", "Dt. pedido",
                          "Dt. prometida", "Sit. pedido", "Tp.doc"), pos, "purchase",
                   {1: "po_number", 2: "po_line", 3: "location", 4: "part_number",
                    5: "supplier_code", 6: "quantity", 7: "unit", 8: "unit_price", 9: "currency",
                    10: "ordered_at", 11: "promised_at", 12: "status"}),
            _Block("Parâmetros de extração", ("Campo", "Valor"),
                   [("Sistema", "ERP industrial - ambiente de teste"),
                    ("Status incluídos", "Encerrado; Em aberto"), ("Data-base", export)]),
        ]),
        _Sheet("Saldos_01", "MM / posição de estoque por centro",
               "Locais produtivos | fotografias por data", [
            _Block(None, ("Dt. posição", "Centro", "Material", "UM", "Físico bruto", "Reservado",
                          "Bloqueado", "Critério saldo"),
                   [(export, plant, part, unit, gross, reserved, blocked,
                     "Livre utilização + reserva + bloqueio")
                    for plant, part, unit, gross, reserved, blocked in _STOCK], "inventory",
                   {1: "as_of", 2: "location", 3: "part_number", 4: "unit", 5: "total",
                    6: "reserved", 7: "blocked"}),
            _Block("Arquivo de posição anterior (consulta)",
                   ("Dt. posição", "Centro", "Material", "UM", "Físico bruto", "Reservado",
                    "Bloqueado", "Origem"),
                   [(old, plant, part, unit, gross, reserved, blocked, "Snapshot anterior")
                    for plant, part, unit, gross, reserved, blocked in _OLD_STOCK], "inventory",
                   {1: "as_of", 2: "location", 3: "part_number", 4: "unit", 5: "total",
                    6: "reserved", 7: "blocked"}),
        ]),
        _Sheet("Ordens", "PP / carteira de ordens e solicitações",
               "Todos os centros de custo | abertas", [
            _Block(None, ("Nº ordem", "Centro", "Área", "Produto / ref.", "Qtd. planejada", "UM",
                          "Necessário em", "Situação", "Origem"),
                   [(ref, plant, area, product, quantity, unit, when(_on(needed)), status, origin)
                    for ref, plant, area, product, quantity, unit, needed, status, origin
                    in _PLANS], "production_plan",
                   {1: "reference", 2: "location", 3: "area", 4: "product_code", 5: "quantity",
                    6: "unit", 7: "need_by", 8: "status"}),
            _Block("Solicitações internas avulsas",
                   ("Req. interna", "Centro", "Área", "Material", "Qtd.", "UM",
                    "Data necessidade"),
                   [(ref, plant, area, part, quantity, unit, when(_on(needed)))
                    for ref, plant, area, part, quantity, unit, needed in _REQUISITIONS],
                   "material_requirement",
                   {1: "reference", 2: "location", 3: "area", 4: "part_number", 5: "quantity",
                    6: "unit", 7: "need_by"}),
        ]),
        _Sheet("Cadastro", "MM / cadastro básico e parceiros",
               "Materiais ativos e fornecedores aprovados", [
            _Block(None, ("Cód. material", "Descrição breve", "UM básica", "Família / área",
                          "Situação", "Classe"),
                   [(part, description, unit, area, "Ativo", "Industrial")
                    for part, description, unit, area in _ITEMS], "item",
                   {1: "part_number", 2: "description", 3: "unit", 4: "area"}),
            _Block("Diretório de parceiros de fornecimento",
                   ("Cód. parceiro", "Razão social", "País", "Cadastro", "Status"),
                   [(s.code, s.name, "BR", when(_on("2026-02-10")), "Ativo") for s in SUPPLIERS],
                   "supplier", {1: "supplier_code", 2: "name", 3: "country"}),
            _Block("Associação material x fonte homologada",
                   ("Material", "Parceiro", "Sit. homologação", "Dt. validação"),
                   [(part, s.code, "Homologado", when(_on(since)))
                    for part, s, since in _APPROVED], "supplier_item",
                   {1: "part_number", 2: "supplier_code", 3: "status", 4: "valid_from"}),
        ]),
        _Sheet("Recebimentos", "MM / movimentação de entrada e agenda",
               "Mov. 101 encerrados + entregas pendentes", [
            _Block(None, ("Pedido", "Item", "Centro", "Dt. entrada", "Qtd. recebida",
                          "Qtd. rejeitada", "UM", "Mov.", "Referência NF"), receipts, "receipt",
                   {1: "po_number", 2: "po_line", 3: "location", 4: "received_at",
                    5: "quantity", 6: "rejected", 7: "unit", 8: "movement", 9: "receipt_id"}),
            _Block("Agenda de confirmações ainda sem entrada física",
                   ("Pedido", "Item", "Centro", "Material", "Qtd. prevista", "UM",
                    "Previsão entrega", "Sit. entrega", "Fonte previsão"), pending,
                   "delivery_schedule",
                   {1: "po_number", 2: "po_line", 3: "location", 4: "part_number", 5: "quantity",
                    6: "unit", 7: "expected_at", 8: "status"}),
        ]),
        _Sheet("Contatos", "Parceiros / contatos comerciais",
               "Contatos em uso - unidades industriais", [
            _Block(None, ("ID parceiro", "Contato / equipe", "E-mail", "Setor", "Horário ref.",
                          "País"),
                   [(s.code, name, s.email, sector, hours, "BR")
                    for s, name, sector, hours in _CONTACTS], "supplier",
                   {1: "supplier_code", 3: "email"}),
        ]),
        _Sheet("Engenharia", "PP / estrutura de fabricação vigente",
               "Estruturas alternativas não incluídas", [
            _Block(None, ("Cód. produto", "Componente", "Consumo por 1", "UM comp.",
                          "Rev. estrutura", "Vigente desde"),
                   [(product, part, per, unit, revision, when(_on(since)))
                    for product, part, per, unit, revision, since in _BOM], "bom",
                   {1: "product_code", 2: "part_number", 3: "quantity_per", 4: "unit",
                    5: "revision", 6: "valid_from"}),
            _Block("Cadastro referência produto acabado",
                   ("Produto", "Designação", "UM produto", "Área responsável"),
                   [*_PRODUCTS], "item",
                   {1: "part_number", 2: "description", 3: "unit", 4: "area"}),
        ]),
        _Sheet("Comercial", "CX / histórico de comunicações de cotação",
               "Arquivo encerrado | referência COT-ANT-087", [
            _Block(None, ("Ref. cotação", "Forn. ID", "Dt./hora comunicação", "Assunto",
                          "Mensagem (corpo e-mail)", "Rev."),
                   [("COT-ANT-087", "F001", when(_on("2026-07-15")), "Cotação histórica",
                     "Para o rolamento cód. 000101 podemos ofertar 60 pcs a R$ 41,80/pcs (BRL), "
                     "entrega em 14 dias corridos após aceite.", "R00"),
                    ("COT-ANT-087", "F003", when(_on("2026-07-16")), "Proposta mista",
                     "Para o eixo 000102 ofertamos 25 pcs a R$ 245,00/pcs (BRL). Oferecemos "
                     "também correia de manutenção cód. 000302: 15 pcs a R$ 77,00/pcs.", "R00")]),
        ]),
    ]  # fmt: skip


@dataclass(frozen=True, slots=True)
class MockWorkbook:
    data: bytes
    sections: Mapping[str, tuple[TableProposal, ...]]


def mock_workbook(export: date) -> MockWorkbook:
    """The equivalent export dated `export`, and the mapping of each of its tables."""
    from openpyxl import Workbook  # pyright: ignore[reportMissingTypeStubs]

    workbook = cast(Any, Workbook())
    workbook.remove(workbook.active)
    sections: dict[str, tuple[TableProposal, ...]] = {}
    shift = _Shift(export)
    for spec in _sheets(shift):
        sheet = workbook.create_sheet(spec.name)
        sheet.append([spec.title])
        sheet.append(["Emissão:", shift.cell(_BASE)])
        sheet.append(["Seleção:", spec.selection])
        found: list[TableProposal] = []
        for block in spec.blocks:
            if block.title is not None:
                sheet.append([block.title])
            sheet.append(list(block.header))
            header_row = sheet.max_row
            for row in block.rows:
                sheet.append(list(row))
            if block.kind is not None and block.rows:
                found.append(
                    TableProposal(
                        kind=block.kind,
                        header_row=header_row,
                        first_row=header_row + 1,
                        last_row=header_row + len(block.rows),
                        columns=[
                            ColumnMapping(field=name, column=column)
                            for column, name in block.fields.items()
                        ],
                    )
                )
        sections[spec.name] = tuple(found)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return MockWorkbook(data=buffer.getvalue(), sections=sections)


def scripted_mapper(sections: Mapping[str, Sequence[TableProposal]]) -> Mapper:
    """What a correct source mapper proposes for each page of the equivalent workbook."""

    async def mapper(task: MappingTask, context: ToolContext) -> MappingProposal:
        return MappingProposal(
            sections=[
                s.model_copy(
                    update={
                        "first_row": max(s.first_row, task.first_row),
                        "last_row": min(s.last_row, task.last_row),
                    }
                )
                for s in sections.get(task.table, ())
                if s.last_row >= task.first_row and s.first_row <= task.last_row
            ]
        )

    return mapper


def mock_case(export: date) -> Case:
    """The equivalent workbook dated `export`, with its golden answer.

    Quantities, suppliers and last prices are the original export's; the reference
    prices are computed from this workbook's own closed orders.
    """
    book = mock_workbook(export)
    parts: list[PartGolden] = []
    for real in REAL_PARTS:
        orders = [o for o in _CLOSED if o.part == real.part]
        target, maximum = reference_prices(orders)
        parts.append(
            PartGolden(
                part=real.part,
                description=real.description,
                unit=real.unit,
                to_buy=real.to_buy,
                last_prices=real.last_prices,
                open_price=real.open_price,
                target=target,
                maximum=maximum,
            )
        )
    return Case(
        name="mock_export_mtr100",
        filename=MOCK_FILENAME,
        data=book.data,
        export=export,
        parts=tuple(parts),
        mapper=scripted_mapper(book.sections),
    )


def supplier(key: str) -> SupplierRef:
    return next(s for s in SUPPLIERS if s.key == key)
