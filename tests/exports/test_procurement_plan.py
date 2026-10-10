"""The full purchase plan: one row per material and plant, the same in the file and the pages."""

import io
from datetime import date
from decimal import Decimal
from typing import Any, cast

from bido_back.exports.procurement_plan import (
    MAX_PAGE,
    page_rows,
    plan_counts,
    plan_filename,
    plan_rows,
    plan_workbook,
)
from bido_back.purchasing.sources.contracts import (
    AffectedDemand,
    AssessmentLine,
    Assumption,
    ExpediteAction,
    PlanningScope,
    PriceBenchmark,
    SourceAssessment,
    SourcingProposal,
    SupplierEvidence,
    TimingCheck,
)

AS_OF = date(2026, 10, 8)


def _supplier(
    code: str, name: str, *, price: str, currency: str, recommended: bool
) -> SupplierEvidence:
    return SupplierEvidence(
        supplier_code=code,
        name=name,
        contract_price=Decimal(price),
        contract_currency=currency,
        estimated_arrival=date(2026, 11, 2),
        recommended=recommended,
        suggested=recommended,
    )


def _line(
    part: str,
    plant: str,
    to_buy: str | None,
    *,
    outcome: str | None = None,
    sourcing: bool = False,
    currency: str = "BRL",
    price: str = "10",
) -> AssessmentLine:
    timing = (
        TimingCheck.model_validate(
            {
                "order_date": AS_OF,
                "outcome": outcome,
                "first_shortage_date": date(2026, 10, 12),
                "first_stockout_date": date(2026, 10, 26),
                "affected_demand": [
                    AffectedDemand(
                        reference="0000004113",
                        product_code="300000030",
                        need_by=date(2026, 10, 12),
                        quantity=Decimal(10),
                    ).model_dump()
                ],
            }
        )
        if outcome is not None
        else None
    )
    return AssessmentLine(
        part_number=part,
        description=f"Material {part}",
        location=plant,
        unit="UN",
        available=Decimal(5),
        safety_stock=Decimal(2),
        on_order=Decimal(1),
        calculated_to_buy=Decimal(to_buy) if to_buy is not None else None,
        net_requirement=Decimal(7) if to_buy == "10" else None,
        suppliers=(
            []
            if sourcing
            else [
                _supplier(
                    "100", f"Supplier {part}", price=price, currency=currency, recommended=True
                )
            ]
        ),
        timing=timing,
        sourcing=SourcingProposal(excluded=["blocked"]) if sourcing else None,
        expedite=(
            [
                ExpediteAction(
                    po_number="4500019736",
                    po_line="10",
                    supplier_name="Motores do Sul",
                    quantity=Decimal(20),
                    current_date=date(2026, 12, 11),
                    requested_date=date(2026, 10, 12),
                    reason="mrp_reschedule_in",
                )
            ]
            if outcome == "unavoidable_stockout"
            else []
        ),
        assumptions=(
            [
                Assumption(
                    code="lot_size_applied",
                    message="The net need of 7 UN was raised to 10 UN by the minimum lot 10.",
                    part_numbers=[part],
                )
            ]
            if to_buy == "10"
            else []
        ),
    )


def _assessment() -> SourceAssessment:
    lines = [
        _line("A", "1000", "10"),  # purchase, lot rule applied
        _line("B", "1000", "5", outcome="unavoidable_stockout"),
        _line("C", "1000", "4", outcome="meets_consuming_safety_stock"),
        _line("D", "2000", "3", sourcing=True),
        _line("E", "2000", "0"),  # covered
        _line("F", "2000", None),  # covered: nothing calculated to buy
        _line("G", "2000", "2", currency="USD", price="3"),
    ]
    return SourceAssessment(
        lines=lines,
        scope=PlanningScope(
            location="1000", as_of=AS_OF, horizon_end=date(2026, 12, 21), all_materials=True
        ),
    )


def test_one_row_per_material_and_plant_with_a_status_each() -> None:
    rows = plan_rows(_assessment())

    assert [(r.part_number, r.plant, r.status) for r in rows] == [
        ("A", "1000", "purchase"),
        ("B", "1000", "unavoidable_stockout"),
        ("C", "1000", "consumes_safety_stock"),
        ("D", "2000", "no_valid_source"),
        ("E", "2000", "covered"),
        ("F", "2000", "covered"),
        ("G", "2000", "purchase"),
    ]
    a, b = rows[0], rows[1]
    # The net shortfall and the quantity after the lot rules are separate columns.
    assert (a.net_shortfall, a.quantity) == (Decimal(7), Decimal(10))
    assert a.lot_rules is not None and "minimum lot" in a.lot_rules
    assert (a.supplier, a.supplier_code, a.unit_price, a.total) == (
        "Supplier A",
        "100",
        Decimal(10),
        Decimal(100),
    )
    assert b.first_shortage == date(2026, 10, 12)
    assert b.affected_orders == ["0000004113 300000030 (2026-10-12)"]
    assert b.actions == ["Antecipar 4500019736/10 (Motores do Sul): 11/12/2026 → 12/10/2026"]
    covered = rows[4]
    assert covered.supplier is None and covered.total is None
    assert rows[3].actions == ["Buscar fornecedor (sem fonte válida)"]


def test_counts_per_status_and_plant_and_totals_never_mix_currencies() -> None:
    counts = plan_counts(plan_rows(_assessment()))

    assert counts.rows == 7
    assert counts.by_status == {
        "purchase": 2,
        "unavoidable_stockout": 1,
        "consumes_safety_stock": 1,
        "no_valid_source": 1,
        "covered": 2,
    }
    assert counts.by_plant == {"1000": 3, "2000": 4}
    assert counts.to_buy_by_plant == {"1000": 3, "2000": 2}
    assert [(t.currency, t.lines, t.total) for t in counts.totals] == [
        ("BRL", 3, Decimal(100) + Decimal(50) + Decimal(40)),
        ("USD", 1, Decimal(6)),
    ]


def test_pages_filter_by_plant_status_and_material_and_count_every_match() -> None:
    rows = plan_rows(_assessment())

    plant = page_rows(rows, plant=" 2000 ", limit=2)
    assert (plant.matched, plant.pages, [r.part_number for r in plant.rows]) == (4, 2, ["D", "E"])
    second = page_rows(rows, plant="2000", limit=2, page=2)
    assert [r.part_number for r in second.rows] == ["F", "G"]
    # A page past the last is the last.
    assert page_rows(rows, plant="2000", limit=2, page=9).page == 2
    stockouts = page_rows(rows, statuses=["unavoidable_stockout", "no_valid_source"])
    assert [r.part_number for r in stockouts.rows] == ["B", "D"]
    assert [r.part_number for r in page_rows(rows, material="material g").rows] == ["G"]
    assert page_rows(rows, limit=10_000).rows == rows[:MAX_PAGE]


def test_a_large_plan_keeps_every_row_in_the_pages() -> None:
    lines = [
        _line(f"M{i:03}", plant, "1") for plant in ("1000", "2000", "3000") for i in range(170)
    ]
    rows = plan_rows(SourceAssessment(lines=lines))

    assert len(rows) == 510
    seen = [
        r.part_number
        for page in range(1, page_rows(rows, plant="2000").pages + 1)
        for r in page_rows(rows, plant="2000", page=page).rows
    ]
    assert len(seen) == 170 == len(set(seen))


def test_the_workbook_holds_a_summary_every_row_a_sheet_per_plant_actions_and_assumptions() -> None:
    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    assessment = _assessment().model_copy(
        update={"assumptions": [Assumption(code="as_of_from_stock", message="Stock dated 08/10.")]}
    )
    workbook = cast(Any, load_workbook(io.BytesIO(plan_workbook(assessment, "pt"))))

    assert workbook.sheetnames == [
        "Resumo",
        "Plano",
        "Centro 1000",
        "Centro 2000",
        "Ações",
        "Premissas",
    ]
    plan = list(workbook["Plano"].iter_rows(values_only=True))
    assert plan[0][:3] == ("Material", "Descrição", "Centro")
    assert len(plan) - 1 == 7
    assert [row[12] for row in plan[1:]] == [
        "Comprar",
        "Ruptura inevitável",
        "Consome estoque de segurança até a chegada",
        "Sem fonte válida",
        "Coberto, sem compra",
        "Coberto, sem compra",
        "Comprar",
    ]
    assert len(list(workbook["Centro 2000"].iter_rows(values_only=True))) - 1 == 4
    summary = {row[0]: row[1] for row in workbook["Resumo"].iter_rows(values_only=True)}
    assert summary["Linhas (material × centro)"] == 7
    assert summary["Total estimado (USD)"] == 6
    actions = list(workbook["Ações"].iter_rows(values_only=True))
    assert [row[0] for row in actions[1:]] == ["B", "D"]
    assumptions = list(workbook["Premissas"].iter_rows(values_only=True))
    assert ("Stock dated 08/10.", None) in assumptions
    english = cast(Any, load_workbook(io.BytesIO(plan_workbook(assessment, "en"))))
    assert english.sheetnames[:2] == ["Summary", "Plan"]


def test_a_formula_like_value_from_the_files_is_written_as_text() -> None:
    line = _line("=HYPERLINK(1)", "1000", "1")
    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    workbook = cast(Any, load_workbook(io.BytesIO(plan_workbook(SourceAssessment(lines=[line])))))
    assert workbook["Plano"]["A2"].value == "'=HYPERLINK(1)"


def test_the_filename_carries_the_reference_date() -> None:
    assert plan_filename(_assessment()) == "plano-de-compras-2026-10-08.xlsx"
    assert plan_filename(SourceAssessment(), "en") == "purchase-plan.xlsx"


def test_price_falls_back_from_an_expired_agreement_to_the_last_price_then_the_target() -> None:
    expired = SupplierEvidence(
        supplier_code="1",
        name="Old",
        contract_price=Decimal(9),
        contract_currency="BRL",
        contract_expired=True,
        last_price=Decimal(12),
        currency="BRL",
        recommended=True,
    )
    with_last = AssessmentLine(part_number="X", calculated_to_buy=Decimal(1), suppliers=[expired])
    target = AssessmentLine(
        part_number="Y",
        calculated_to_buy=Decimal(1),
        price=PriceBenchmark.model_validate(
            {
                "currency": "BRL",
                "basis": "received_orders",
                "observations": 3,
                "window_days": 365,
                "last_price": "11",
                "avg_price": "11",
                "min_price": "10",
                "target_price": "10.5",
                "max_price": "12",
            }
        ),
    )
    rows = plan_rows(SourceAssessment(lines=[with_last, target]))
    assert (rows[0].unit_price, rows[1].unit_price) == (Decimal(12), Decimal("10.5"))


def test_totals_are_in_cents_and_a_price_per_thousand_keeps_six_decimals() -> None:
    line = _line("S", "1000", "3", price="0.30346666666")
    row = plan_rows(SourceAssessment(lines=[line]))[0]
    assert row.unit_price == Decimal("0.303467")
    assert row.total == Decimal("0.91")
