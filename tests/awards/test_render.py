"""The purchase-order document and email: pure code, no database."""

import io
from datetime import date
from decimal import Decimal
from typing import Any, cast

from bido_back.awards.render import (
    PoDocument,
    PoLine,
    content_hash,
    line_total,
    po_email_body,
    to_html,
    to_xlsx,
)


def doc(price: str = "0.10", description: str = "MLCC 100nF") -> PoDocument:
    quantity = Decimal(4000)
    total = line_total(quantity, Decimal(price))
    return PoDocument(
        ref="PO-7",
        rfq_ref="RFQ-3",
        issued_on=date(2026, 10, 7),
        buyer_company="Voltwerk GmbH",
        supplier_name="Rhein <Elektronik>",
        currency="EUR",
        incoterm="DAP Leipzig",
        payment_terms="30 days net",
        total=total,
        lines=(
            PoLine(
                position=1,
                description=description,
                mpn="GRM188R71H104KA93D",
                quantity=quantity,
                unit="pcs",
                unit_price=Decimal(price),
                line_total=total,
                need_by=date(2027, 1, 15),
                lead_time_days=42,
            ),
        ),
    )


def test_line_totals_round_to_cents() -> None:
    assert line_total(Decimal(3), Decimal("0.3333")) == Decimal("1.00")


def test_the_hash_changes_with_what_the_supplier_receives() -> None:
    assert content_hash(doc()) == content_hash(doc())
    assert content_hash(doc()) != content_hash(doc(price="0.11"))
    assert content_hash(doc()) != content_hash(doc(description="MLCC 220nF"))


def test_html_escapes_supplier_text() -> None:
    page = to_html(doc(description="<script>x</script>"))
    assert "<script>x" not in page and "&lt;script&gt;x" in page
    assert "Rhein &lt;Elektronik&gt;" in page and "400.00" in page


def test_xlsx_keeps_numbers_and_neutralizes_formulas() -> None:
    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    data = to_xlsx(doc(description="=HYPERLINK(1)"))
    sheet = cast(Any, load_workbook(io.BytesIO(data))).active
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    header = rows.index(
        [
            "Pos",
            "Description",
            "MPN",
            "Quantity",
            "Unit",
            "Unit price",
            "Currency",
            "Line total",
            "Need by",
            "Lead time (days)",
            "Part number",
        ]
    )
    line = rows[header + 1]
    assert line[1] == "'=HYPERLINK(1)"
    assert line[3] == 4000 and line[7] == 400


def test_the_email_is_in_the_supplier_language_with_an_english_fallback() -> None:
    german = po_email_body(doc(), "de", "Anna Becker")
    assert "Bestellung PO-7" in german and "RFQ-3" in german
    assert "4000 pcs x 0.1 EUR = 400.00 EUR" in german
    assert german.endswith("Anna Becker\nVoltwerk GmbH")
    assert "purchase order PO-7" in po_email_body(doc(), "zh", "Anna")


def test_the_po_email_signs_the_company_once_without_a_buyer_name() -> None:
    # Production: the buyer's name falls back to the company (email audit F4).
    body = po_email_body(doc(), "pt", "Voltwerk GmbH")
    assert body.endswith("Atenciosamente,\nVoltwerk GmbH")
    assert body.count("Voltwerk GmbH") == 1
