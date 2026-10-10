"""A large SAP export across three plants is planned as the reference plan says, by code.

The scenario (`evals.sap_large`) carries the traps of the CEO's large-scale test: open
orders with several schedule lines ("4500102721/00010/0001" in MD04, item 10 in ME2N),
overdue lines, per-plant safety stock and lot rules, per-1,000 prices, 18 L buckets,
supplier minimum orders, expired info records, one company with a blocked and an active
supplier code, and consumption-based (ND) materials. The expected plan is computed from
the generated facts by an independent reference of the CEO's rules.
"""

import io
import logging
from datetime import date
from decimal import Decimal
from uuid import uuid4

import pytest

from bido_back.agents.context import ToolContext
from bido_back.purchasing.sources.contracts import (
    AssessmentLine,
    ColumnMapping,
    MappingProposal,
    MappingTask,
    SourceAssessment,
    TableProposal,
)
from bido_back.purchasing.sources.joins import state
from bido_back.purchasing.sources.mapping import Mapper
from bido_back.purchasing.sources.packs import read_known
from bido_back.purchasing.sources.packs.base import PackContext
from bido_back.purchasing.sources.readers import SourceDocument, read_source
from bido_back.purchasing.sources.service import analyze_documents
from evals.sap_large import Scenario, generate, reports
from evals.sap_large_check import assess, compare

MATERIALS = 30
SEED = 7


def _plan(result: SourceAssessment) -> dict[tuple[str, str], tuple[Decimal | None, str | None]]:
    return {
        (line.part_number, line.location or ""): (
            line.calculated_to_buy,
            next((s.supplier_code for s in line.suppliers if s.recommended), None),
        )
        for line in result.lines
    }


def _line(result: SourceAssessment, part: str, plant: str) -> AssessmentLine:
    return next(x for x in result.lines if x.part_number == part and x.location == plant)


async def test_three_plants_are_planned_as_the_reference_says() -> None:
    scenario = generate(MATERIALS, SEED)
    result = await assess(scenario)
    report = compare(scenario, result)
    assert report.mismatches == []
    # Every kind of outcome is exercised, none of the forbidden ones happens.
    for kind in ("buy", "covered", "no_valid_source"):
        assert report.counts[f"got_{kind}"] == report.counts[f"want_{kind}"] > 0, kind
    assert report.counts["got_unavoidable"] == report.counts["want_unavoidable"] > 0
    assert report.counts["nd_excluded"] > 0 and report.counts["nd_planned"] == 0
    assert report.counts["blocked_recommended"] == 0
    assert report.counts["expired_recommended"] == 0
    assert report.counts["open_orders_not_netted"] == 0
    excluded = next(a for a in result.assumptions if a.code == "consumption_based_excluded")
    assert "MRP type ND" in excluded.message


async def test_every_schedule_line_of_an_open_order_is_netted() -> None:
    scenario = generate(MATERIALS, SEED)
    result = await assess(scenario)
    pair = next(p for p in scenario.pairs if len(p.lines) == 3 and p.mrp_type == "PD")
    line = _line(result, pair.material, pair.plant)
    # MD04 numbers the schedule lines ("/0002"); the order list's rows take its numbers.
    assert sorted((int(i.schedule_line or 0), i.quantity) for i in line.inbound) == sorted(
        (s.line, Decimal(s.open_quantity)) for s in pair.lines
    )
    assert not [a for a in line.assumptions if a.code == "open_orders_not_counted"]


async def test_the_safety_stock_and_lot_rules_are_the_plants_own() -> None:
    scenario = generate(MATERIALS, SEED)
    result = await assess(scenario)
    by_material: dict[str, list[int]] = {}
    for p in scenario.pairs:
        if p.mrp_type == "PD":
            by_material.setdefault(p.material, []).append(p.safety)
    material = next(m for m, values in by_material.items() if len(set(values)) == 3)
    for pair in (p for p in scenario.pairs if p.material == material and p.mrp_type == "PD"):
        line = _line(result, pair.material, pair.plant)
        assert line.safety_stock == Decimal(pair.safety)
        if line.calculated_to_buy and line.net_requirement is not None:
            # The net need and the order quantity are both shown.
            assert line.calculated_to_buy >= line.net_requirement


async def test_the_duplicated_company_is_judged_by_supplier_code() -> None:
    scenario = generate(MATERIALS, SEED)
    result = await assess(scenario)
    recommended = {
        s.supplier_code for line in result.lines for s in line.suppliers if s.recommended
    }
    assert "100151" in recommended  # the active code of the company
    assert "100150" not in recommended  # its purchase-blocked code
    assert "100140" not in recommended


async def test_planned_orders_do_not_change_the_plan_when_the_mrp_list_is_there() -> None:
    scenario = generate(MATERIALS, SEED)
    without = await assess(scenario)
    with_md16 = await assess(scenario, planned_orders=True)
    assert "MD16" in {r.report for r in with_md16.reports}
    assert _plan(with_md16) == _plan(without)
    assert "planned_orders_in_mrp_list" in {a.code for a in with_md16.assumptions}
    # The planned orders name the products of the orders a stockout puts at risk.
    at_risk = [
        demand
        for line in with_md16.lines
        if line.timing is not None
        for demand in line.timing.affected_demand
    ]
    assert at_risk and all(
        d.product_description is not None and d.product_description.startswith("PRODUTO")
        for d in at_risk
    )


def test_planned_orders_alone_are_a_production_plan_of_planned_orders() -> None:
    scenario = generate(4, SEED)
    data = reports(scenario, planned_orders=True)["MD16.xlsx"]
    tables, _ = read_source(SourceDocument("md16", "MD16.xlsx", "application/vnd.ms-excel", data))
    known = read_known(tables, PackContext(filenames={}, today=date(2026, 10, 9)))
    assert [r.report for r in known.reports] == ["MD16"]
    plans = [r for r in known.records if r.kind == "production_plan"]
    assert plans and all(state(r) == "planned" for r in plans)  # firmed or not, never unknown
    assert "planned_orders_exploded" in {a.code for a in known.assumptions}


def _extra_file(rows: list[list[object]]) -> bytes:
    from openpyxl import Workbook  # pyright: ignore[reportMissingTypeStubs]

    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _mapped_as(kind: str, fields: list[str]) -> Mapper:
    """A model mapper that reads the unknown table's columns as `kind` records."""

    async def mapper(task: MappingTask, _: ToolContext) -> MappingProposal:
        return MappingProposal(
            sections=[
                TableProposal(
                    kind=kind,  # pyright: ignore[reportArgumentType]
                    header_row=1,
                    first_row=max(2, task.first_row),
                    last_row=task.last_row,
                    columns=[ColumnMapping(field=f, column=i) for i, f in enumerate(fields, 1)],
                )
            ]
        )

    return mapper


async def _with_extra(scenario: Scenario, extra: bytes, mapper: Mapper) -> SourceAssessment:
    documents = [
        SourceDocument(name, name, "application/vnd.ms-excel", data)
        for name, data in reports(scenario).items()
    ]
    documents.append(SourceDocument("extra", "custom.xlsx", "application/vnd.ms-excel", extra))
    return await analyze_documents(
        documents,
        "O que comprar, quanto e de quem",
        ToolContext(tenant_id=uuid4(), user_id=uuid4(), run_id=uuid4()),
        mapper,
        today=date(2026, 10, 9),
    )


async def test_an_unknown_file_read_by_the_model_never_switches_or_empties_the_plan(
    caplog: pytest.LogCaptureFixture,
) -> None:
    scenario = generate(MATERIALS, SEED)
    baseline = _plan(await assess(scenario))
    plant = scenario.pairs[0].plant
    product = "300000000"
    # Planned orders in an unrecognized layout: read as production plans, they would turn
    # the plan into a product build with unknown order states and no lines (the MD16 case).
    orders = _extra_file(
        [
            ["Nº da ordem", "Produto", "Planta", "Qtd", "Fim", "Situação"],
            *[[f"88{n:04d}", product, plant, 50, date(2026, 11, n + 1), "LA"] for n in range(5)],
        ]
    )
    caplog.set_level(logging.INFO, logger="bido_back.purchasing.sources.mapping")
    switched = await _with_extra(
        scenario,
        orders,
        _mapped_as(
            "production_plan",
            ["reference", "product_code", "location", "quantity", "need_by", "status"],
        ),
    )
    assert _plan(switched) == baseline
    assert "mapped_mode_records_left_out" in {i.code for i in switched.issues}
    # Stock of a planned item in another layout repeats MB52: the report's stock is used.
    pair = next(p for p in scenario.pairs if p.mrp_type == "PD")
    stock = _extra_file(
        [
            ["Item", "Planta", "Saldo", "Unid"],
            [pair.material, pair.plant, 999_999, pair.unit],
        ]
    )
    repeated = await _with_extra(
        scenario,
        stock,
        _mapped_as("inventory", ["part_number", "location", "available", "unit"]),
    )
    assert _plan(repeated) == baseline
    assert "mapped_records_overlap_reports" in {i.code for i in repeated.issues}
    # One structured line per unrecognized table: headers, kinds, rows; no values.
    logged = [r for r in caplog.records if r.getMessage() == "source_unrecognized_table"]
    assert logged
    first = logged[0]
    assert first.__dict__["headers"] == [
        "no da ordem",
        "produto",
        "planta",
        "qtd",
        "fim",
        "situacao",
    ]
    assert first.__dict__["kinds"] == ["production_plan"]
    assert first.__dict__["rows"] == 6
    assert product not in str(first.__dict__)
