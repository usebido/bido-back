"""The saved evidence API uses authenticated identity and server-created assessments."""

from uuid import uuid4

import httpx
import pytest
from fastapi import FastAPI
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession, async_sessionmaker

from bido_back.api.deps import get_db, get_principal, get_storage
from bido_back.api.procurement_assessments import get_source_mapper, procurement_assessments_router
from bido_back.auth import Principal
from bido_back.db import create_sessionmaker
from bido_back.purchasing.sources.mapping import Mapper
from bido_back.storage import BlobStorage
from tests.approvals.conftest import People
from tests.copilot.test_source_tools import SourceMapper, TrackingStorage
from tests.purchasing.test_api import build_app
from tests.purchasing.test_source_snapshots import upload_source
from tests.rfqs.test_service import buyer


def _app() -> FastAPI:
    app = build_app()
    app.include_router(procurement_assessments_router())
    return app


async def test_assessment_routes_require_authentication() -> None:
    transport = httpx.ASGITransport(app=_app())
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as http:
        responses = [
            await http.get("/v1/procurement-assessments"),
            await http.post(
                "/v1/procurement-assessments",
                json={
                    "upload_ids": [str(uuid4())],
                    "objective": "Analyze material requirements",
                },
            ),
            await http.get(f"/v1/procurement-assessments/{uuid4()}"),
            await http.post(
                f"/v1/procurement-assessments/{uuid4()}/rfq",
                json={
                    "title": "Components",
                    "base_currency": "EUR",
                },
            ),
        ]
    assert [r.status_code for r in responses] == [401, 401, 401, 401]


@pytest.mark.db
async def test_assessment_http_roundtrip_and_owner_authorization(
    app_engine: AsyncEngine,
    purchasing_rows: People,
) -> None:
    app = _app()
    who = buyer(purchasing_rows)
    identity = {"principal": who}
    storage, mapper = TrackingStorage(), SourceMapper()
    upload_id = await upload_source(app_engine, who, storage)
    sessionmaker = create_sessionmaker(app_engine)

    def principal() -> Principal:
        return identity["principal"]

    def db() -> async_sessionmaker[AsyncSession]:
        return sessionmaker

    def blobs() -> BlobStorage:
        return storage

    def source_mapper() -> Mapper:
        return mapper

    app.dependency_overrides[get_principal] = principal
    app.dependency_overrides[get_db] = db
    app.dependency_overrides[get_storage] = blobs
    app.dependency_overrides[get_source_mapper] = source_mapper
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as http:
        request = {"upload_ids": [str(upload_id)], "objective": "Analyze material requirements"}
        created = await http.post("/v1/procurement-assessments", json=request)
        assert created.status_code == 200, created.text
        assessment = created.json()
        assert assessment["persisted"] is True
        path = f"/v1/procurement-assessments/{assessment['assessment_id']}"
        listing = await http.get("/v1/procurement-assessments", params={"search": "MATERIAL"})
        assert listing.status_code == 200
        assert [row["id"] for row in listing.json()] == [assessment["assessment_id"]]
        assert "lines" not in listing.json()[0]
        assert (
            await http.get("/v1/procurement-assessments", params={"limit": 51})
        ).status_code == 422
        fetched = await http.get(path)
        assert fetched.json() == assessment
        assert (await http.post("/v1/procurement-assessments", json=request)).json() == assessment
        forged = await http.post(
            "/v1/procurement-assessments",
            json={
                **request,
                "assessment": {"status": "assessed"},
            },
        )
        assert forged.status_code == 422
        unsupported = await http.post(
            f"{path}/rfq", json={"title": "Components", "base_currency": "XYZ"}
        )
        assert unsupported.status_code == 422
        assert (await http.get(path)).json()["rfq_id"] is None
        rfq_request = {"title": "Components", "base_currency": "EUR"}
        draft = await http.post(f"{path}/rfq", json={**rfq_request, "base_currency": " eur "})
        assert draft.status_code == 200, draft.text
        again = await http.post(f"{path}/rfq", json=rfq_request)
        assert again.json()["rfq_id"] == draft.json()["rfq_id"]
        assert again.json()["reused"] is True
        for foreign in (
            Principal(tenant_id=who.tenant_id, user_id=purchasing_rows.approver, role="approver"),
            Principal(
                tenant_id=purchasing_rows.other_tenant,
                user_id=purchasing_rows.other_approver,
                role="approver",
            ),
        ):
            identity["principal"] = foreign
            assert (await http.get("/v1/procurement-assessments")).json() == []
            assert (await http.get(path)).status_code == 404
            assert (await http.post(f"{path}/rfq", json=rfq_request)).status_code == 404
            assert (await http.post("/v1/procurement-assessments", json=request)).status_code == 404


@pytest.mark.db
async def test_the_full_plan_downloads_as_xlsx_for_its_owner_only(
    app_engine: AsyncEngine,
    purchasing_rows: People,
) -> None:
    from io import BytesIO
    from typing import Any, cast

    from openpyxl import load_workbook  # pyright: ignore[reportMissingTypeStubs]

    app = _app()
    who = buyer(purchasing_rows)
    identity = {"principal": who}
    storage, mapper = TrackingStorage(), SourceMapper()
    upload_id = await upload_source(app_engine, who, storage)
    sessionmaker = create_sessionmaker(app_engine)
    app.dependency_overrides[get_principal] = lambda: identity["principal"]
    app.dependency_overrides[get_db] = lambda: sessionmaker
    app.dependency_overrides[get_storage] = lambda: storage
    app.dependency_overrides[get_source_mapper] = lambda: mapper
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as http:
        created = await http.post(
            "/v1/procurement-assessments",
            json={"upload_ids": [str(upload_id)], "objective": "Analyze material requirements"},
        )
        assessment = created.json()
        path = f"/v1/procurement-assessments/{assessment['assessment_id']}/plan.xlsx"
        response = await http.get(path, params={"lang": "en"})
        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "attachment; filename=" in response.headers["content-disposition"]
        workbook = cast(Any, load_workbook(BytesIO(response.content)))
        assert workbook.sheetnames[:2] == ["Summary", "Plan"]
        plan_rows = list(workbook["Plan"].iter_rows(values_only=True))
        assert len(plan_rows) - 1 == len(assessment["lines"])
        assert (await http.get(path, params={"lang": "xx"})).status_code == 422
        identity["principal"] = Principal(
            tenant_id=purchasing_rows.other_tenant,
            user_id=purchasing_rows.other_approver,
            role="approver",
        )
        assert (await http.get(path)).status_code == 404
    unauthenticated = httpx.ASGITransport(app=_app())
    async with httpx.AsyncClient(transport=unauthenticated, base_url="http://test") as http:
        assert (await http.get(path)).status_code == 401
