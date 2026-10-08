"""Lossless source windows, independent of supported business import types."""

import io
from datetime import date
from typing import Any, cast

import pytest

from bido_back.purchasing.inspection import SpreadsheetWindow, inspect_spreadsheet
from bido_back.purchasing.workbook import WorkbookFileError, read_workbook
from tests.purchasing.test_workbook import workbook


def test_inspection_reads_every_sheet_including_hidden_and_late_wide_cells() -> None:
    from openpyxl import Workbook  # pyright: ignore[reportMissingTypeStubs]

    book = cast(Any, Workbook())
    book.active.title = "Stock"
    text = "A complete long material description " * 5
    book.active.cell(42, 35, text)
    for i in range(20):
        sheet = book.create_sheet(f"Other {i}")
        sheet.cell(1, 1, "000010")
    book.worksheets[-1].sheet_state = "hidden"
    buffer = io.BytesIO()
    book.save(buffer)
    data = buffer.getvalue()

    summary = inspect_spreadsheet(data, "export.xlsx", SpreadsheetWindow())
    assert len(summary.sheets) == 21
    assert summary.sheets[-1].state == "hidden"
    assert summary.cells == []
    page = inspect_spreadsheet(data, "export.xlsx", SpreadsheetWindow(sheet="Stock"))
    assert page.next_column == 21 and page.next_row is None
    last = inspect_spreadsheet(
        data, "export.xlsx", SpreadsheetWindow(sheet="Stock", first_column=21)
    )
    [cell] = last.cells
    assert (cell.row, cell.column, cell.value) == (42, 35, text)
    assert last.next_column is None
    with pytest.raises(WorkbookFileError, match="limited to 20 sheets"):
        read_workbook(data, "export.xlsx")


def test_formulas_codes_formats_and_dates_remain_distinct() -> None:
    from openpyxl import Workbook  # pyright: ignore[reportMissingTypeStubs]

    book = cast(Any, Workbook())
    sheet = book.active
    sheet.title = "Facts"
    sheet.append(["000010", 12.5, date(2026, 10, 31), "=B1*2", "#DIV/0!", False])
    sheet["B1"].number_format = "0.000"
    buffer = io.BytesIO()
    book.save(buffer)
    result = inspect_spreadsheet(buffer.getvalue(), "facts.xlsx", SpreadsheetWindow(sheet="Facts"))
    code, number, day, formula, error, flag = result.cells
    assert code.kind == "text" and code.value == "000010"
    assert number.kind == "number" and number.value == "12.5"
    assert number.number_format == "0.000"
    assert day.kind == "date" and day.value == "2026-10-31T00:00:00"
    assert formula.formula == "=B1*2" and formula.value is None
    assert error.kind == "error" and error.value == "#DIV/0!"
    assert flag.kind == "boolean" and flag.value == "false"


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_paging_covers_source_cells_once_without_changing_values(extension: str) -> None:
    rows: list[list[object]] = [[f"code-{n}", n] for n in range(121)]
    filename = f"source.{extension}"
    data = (
        "\n".join(f"code-{n};{n}" for n in range(121)).encode()
        if extension == "csv"
        else workbook({"Rows": rows})
    )
    sheet = filename if extension == "csv" else "Rows"
    first = 1
    found: list[tuple[int, str | None]] = []
    while True:
        page = inspect_spreadsheet(data, filename, SpreadsheetWindow(sheet=sheet, first_row=first))
        found.extend((c.row, c.value) for c in page.cells if c.column == 1)
        if page.next_row is None:
            break
        first = page.next_row
    assert found == [(n + 1, f"code-{n}") for n in range(121)]


def test_invalid_windows_fail_explicitly() -> None:
    data = workbook({"Data": [["value"]]})
    with pytest.raises(WorkbookFileError, match="sheet not found"):
        inspect_spreadsheet(data, "x.xlsx", SpreadsheetWindow(sheet="Missing"))
    with pytest.raises(WorkbookFileError, match="outside the sheet"):
        inspect_spreadsheet(data, "x.xlsx", SpreadsheetWindow(sheet="Data", first_row=2))
    with pytest.raises(WorkbookFileError, match="not a readable"):
        inspect_spreadsheet(b"not an XLSX", "x.xlsx", SpreadsheetWindow())
